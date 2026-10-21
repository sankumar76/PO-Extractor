"""
Extract Purchase Order fields from PDFs into a single Excel file.

Folder layout expected:
    po_extractor/
    ├── extract_po_to_excel.py   <- this script
    ├── input_pdfs/              <- drop all your PO PDF files here
    └── output/
        └── po_extract.xlsx      <- created automatically

Columns produced (one row per line item):
    Company Name | PO No | PO Date | Item Description | Qty | Rate | Disc Amount | Amount
"""

import re
import os
import pdfplumber
import openpyxl
from openpyxl.styles import Font

INPUT_FOLDER = "input_pdfs"
OUTPUT_FILE = os.path.join("output", "po_extract.xlsx")

# --- Regex patterns tuned to this PO layout ---
PO_NO_PATTERN = re.compile(r"PO No\s*:\s*(\S+)")
PO_DATE_PATTERN = re.compile(r"PO Date\s*:\s*(\S+)")
# Matches lines like: 1  BEASHAFT-BEARING SHAFT -1-(25.4mm)  73181500  200.00 NOS  85.000  0.00  17000.00
ITEM_LINE_PATTERN = re.compile(
    r"^\s*\d+\s+(.+?)\s+\d{6,8}\s+([\d.]+)\s+\w+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$"
)


def extract_company_name(text):
    """Company name is the first line containing 'Limited' (or 'Ltd')."""
    for line in text.splitlines():
        line = line.strip()
        if "Limited" in line or "Ltd" in line:
            return line
    return ""


def extract_fields_from_pdf(pdf_path):
    """Return a list of row dicts, one per line item found in the PDF."""
    with pdfplumber.open(pdf_path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    company_name = extract_company_name(full_text)

    po_no_match = PO_NO_PATTERN.search(full_text)
    po_date_match = PO_DATE_PATTERN.search(full_text)
    po_no = po_no_match.group(1) if po_no_match else ""
    po_date = po_date_match.group(1) if po_date_match else ""

    rows = []
    for line in full_text.splitlines():
        match = ITEM_LINE_PATTERN.match(line)
        if match:
            description, qty, rate, disc_amt, amount = match.groups()
            rows.append({
                "Company Name": company_name,
                "PO No": po_no,
                "PO Date": po_date,
                "Item Description": description.strip(),
                "Qty": float(qty),
                "Rate": float(rate),
                "Disc Amount": float(disc_amt),
                "Amount": float(amount),
            })
    return rows


def main():
    if not os.path.isdir(INPUT_FOLDER):
        print(f"Folder '{INPUT_FOLDER}' not found. Create it and add your PDFs.")
        return

    pdf_files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith(".pdf")]
    if not pdf_files:
        print(f"No PDF files found in '{INPUT_FOLDER}'.")
        return

    all_rows = []
    for filename in pdf_files:
        pdf_path = os.path.join(INPUT_FOLDER, filename)
        try:
            rows = extract_fields_from_pdf(pdf_path)
            if not rows:
                print(f"⚠  No item lines matched in: {filename} (check its layout)")
            all_rows.extend(rows)
            print(f"✔ Processed: {filename} ({len(rows)} item rows)")
        except Exception as e:
            print(f"✘ Failed on {filename}: {e}")

    if not all_rows:
        print("Nothing extracted. No Excel file was created.")
        return

    os.makedirs("output", exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO Extract"

    headers = list(all_rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    for row in all_rows:
        ws.append([row[h] for h in headers])
        for cell in ws[ws.max_row]:
            cell.font = Font(name="Arial")

    widths = [30, 16, 12, 38, 10, 10, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(OUTPUT_FILE)
    print(f"\nDone. Saved {len(all_rows)} rows to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
