"""
PO Extractor - desktop app version (v4).

What's new in this version:
  - Output formatting (fonts, header fill color, borders, column widths)
    matches the customer-provided template exactly.
  - If the output workbook already exists, new PDFs are APPENDED to it
    rather than overwriting it (S.No numbering continues on from where
    it left off).
  - Each customer gets its own sheet/table within the workbook.
  - A cleaner, more professional-looking interface.

Supports TWO vendor PO formats, auto-detected per PDF:
  FORMAT A - "Accucut"-style   (labels: "PO No", "PO Date")
  FORMAT B - "Bull Agro"-style (labels: "Order No.", "Order Date")
      Order Date -> written into the "PO Date" column
      Order No.  -> written into the "PO No" column
      Description is split on the first hyphen "-"

Output columns (per customer sheet):
    S.No | PO Date | PO No | Customer Name | Item Code | Item Description |
    PO Qty | Rate | Dispatch Qty | Balance Qty
"""

import re
import os
import sys
import threading
import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, scrolledtext, messagebox

import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

APP_TITLE = "PO Data Extractor"
APP_SUBTITLE = "Automated Purchase Order to Excel Conversion"

# ============================================================
# Extraction logic
# ============================================================

# ---------- Format A ("PO No" / "PO Date") ----------
PO_NO_PATTERN = re.compile(r"PO No\s*:\s*(\S+)")
PO_DATE_PATTERN = re.compile(r"PO Date\s*:\s*(\S+)")
ITEM_LINE_PATTERN = re.compile(
    r"^\s*\d+\s+(.+?)\s+(?:\d{6,8}\s+)?([\d.]+)\s+\w+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$"
)

# ---------- Format B ("Order No." / "Order Date") ----------
ORDER_NO_PATTERN = re.compile(r"Order No\.?\s*:\s*(.+)")
ORDER_DATE_PATTERN = re.compile(r"Order Date\s*:\s*(.+)")

LEADING_NON_ALPHA_PATTERN = re.compile(r"^([^A-Za-z]*)(.*)$", re.DOTALL)
DATE_PATTERN = re.compile(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{4})")


def normalize_date(raw_date):
    if not raw_date:
        return raw_date
    match = DATE_PATTERN.search(raw_date)
    if not match:
        return raw_date
    day, month, year = match.groups()
    return f"{int(day):02d}/{int(month):02d}/{year}"


def extract_company_name(text):
    for line in text.splitlines():
        stripped = line.strip()
        if "Limited" in stripped or "Ltd" in stripped:
            return stripped
    for line in text.splitlines():
        stripped = line.strip()
        if stripped and stripped.isupper() and not any(ch.isdigit() for ch in stripped):
            return stripped
    return ""


def split_description(description):
    """FORMAT A: split before the first alphabet character."""
    match = LEADING_NON_ALPHA_PATTERN.match(description)
    prefix = (match.group(1) or "").strip()
    remainder = (match.group(2) or "").strip()
    if prefix == "":
        return "", description.strip()
    return prefix, remainder


def split_description_hyphen(description):
    """FORMAT B: split on the first hyphen."""
    stripped = description.strip()
    if "-" in stripped:
        idx = stripped.index("-")
        return stripped[:idx].strip(), stripped[idx + 1:].strip()
    return "", stripped


def find_col_index(header_row, target):
    for i, cell in enumerate(header_row):
        if cell and target.lower() in cell.replace("\n", " ").strip().lower():
            return i
    return None


def detect_format(full_text):
    if re.search(r"PO No\s*:", full_text):
        return "A"
    if re.search(r"Order No\.?\s*:", full_text):
        return "B"
    return None


def extract_format_a(full_text):
    po_no_match = PO_NO_PATTERN.search(full_text)
    po_date_match = PO_DATE_PATTERN.search(full_text)
    po_no = po_no_match.group(1) if po_no_match else ""
    po_date = po_date_match.group(1) if po_date_match else ""

    items = []
    for line in full_text.splitlines():
        match = ITEM_LINE_PATTERN.match(line)
        if match:
            description, qty, rate, disc_amt, amount = match.groups()
            items.append((description.strip(), float(qty), float(rate)))
    return po_no, po_date, items


def extract_format_b(full_text, tables):
    order_no_match = ORDER_NO_PATTERN.search(full_text)
    order_date_match = ORDER_DATE_PATTERN.search(full_text)
    order_no = order_no_match.group(1).strip() if order_no_match else ""
    order_date = order_date_match.group(1).strip() if order_date_match else ""

    items = []
    for table in tables:
        if not table or len(table) < 3:
            continue
        header_row = table[0]
        if not header_row or not any(c and "description" in str(c).lower() for c in header_row):
            continue

        desc_idx = find_col_index(header_row, "description")
        qty_idx = find_col_index(header_row, "ord. qty") or find_col_index(header_row, "qty")
        rate_idx = find_col_index(header_row, "unit cost")
        if desc_idx is None or qty_idx is None or rate_idx is None:
            continue

        for row in table[2:]:
            sl = row[0]
            if sl and str(sl).strip().isdigit():
                raw_desc = row[desc_idx] or ""
                clean_desc = re.sub(r"HSN/SAC.*", "", raw_desc, flags=re.IGNORECASE)
                clean_desc = " ".join(clean_desc.split())
                qty_raw = (row[qty_idx] or "0").replace(",", "")
                rate_raw = (row[rate_idx] or "0").replace(",", "")
                try:
                    qty = float(qty_raw)
                    rate = float(rate_raw)
                except ValueError:
                    continue
                items.append((clean_desc, qty, rate))
    return order_no, order_date, items


def extract_fields_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        fmt = detect_format(full_text)

        if fmt == "A":
            po_no, po_date, items = extract_format_a(full_text)
            split_func = split_description
        elif fmt == "B":
            tables = []
            for page in pdf.pages:
                tables.extend(page.extract_tables())
            order_no, order_date, items = extract_format_b(full_text, tables)
            po_date, po_no = order_date, order_no
            split_func = split_description_hyphen
        else:
            raise ValueError("Unrecognized PO format (no 'PO No'/'PO Date' or "
                              "'Order No.'/'Order Date' labels found)")

    company_name = extract_company_name(full_text)
    po_date = normalize_date(po_date)

    rows = []
    for description, qty, rate in items:
        prefix, remainder = split_func(description)
        rows.append({
            "PO Date": po_date,
            "PO No": po_no,
            "Customer Name": company_name,
            "Item Code": prefix,
            "Item Description": remainder,
            "PO Qty": qty,
            "Rate": rate,
        })
    return rows


# ============================================================
# Excel output: template-matched formatting, per-customer sheets,
# append-to-existing-file support
# ============================================================

HEADERS = ["S.No", "PO Date", "PO No", "Customer Name", "Item Code",
           "Item Description", "PO Qty", "Rate", "Dispatch Qty", "Balance Qty"]

HEADER_FONT = Font(name="Arial", size=11, bold=True)
DATA_FONT = Font(name="Arial", size=11)
HEADER_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
MERGE_COLS = [1, 2, 3, 4]  # S.No, PO Date, PO No, Customer Name

COLUMN_WIDTHS = {1: 6, 2: 12.78, 3: 27.78, 4: 30.78, 5: 29.78, 6: 40.78, 7: 10.78, 9: 14.78}

INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def sanitize_sheet_name(name):
    cleaned = INVALID_SHEET_CHARS.sub("", name or "").strip()
    return (cleaned[:31] if cleaned else "Customer")


def create_sheet_with_header(wb, sheet_name):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    for col_idx, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return ws


def get_or_create_sheet(wb, customer_name):
    name = sanitize_sheet_name(customer_name)
    if name in wb.sheetnames:
        return wb[name]
    return create_sheet_with_header(wb, name)


def get_next_sno(ws):
    if ws.max_row <= 1:
        return 1
    for r in range(ws.max_row, 1, -1):
        val = ws.cell(row=r, column=1).value
        if val is not None:
            try:
                return int(val) + 1
            except (TypeError, ValueError):
                continue
    return 1


def write_po_group(ws, rows):
    """Append one PDF's rows (one PO, possibly several item rows) as a
    single group: S.No/PO Date/PO No/Customer Name centered and merged."""
    start_row = ws.max_row + 1
    s_no = get_next_sno(ws)

    for offset, row in enumerate(rows):
        r = start_row + offset
        values = [s_no, row["PO Date"], row["PO No"], row["Customer Name"],
                  row["Item Code"], row["Item Description"], row["PO Qty"],
                  row["Rate"], None, None]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx in MERGE_COLS:
                cell.alignment = CENTER_ALIGN

    end_row = start_row + len(rows) - 1
    if end_row > start_row:
        for col in MERGE_COLS:
            ws.merge_cells(start_row=start_row, start_column=col,
                            end_row=end_row, end_column=col)


def build_or_append_workbook(output_path, rows_per_pdf):
    """rows_per_pdf: list of (list_of_row_dicts) - one entry per successfully
    processed PDF, in the order they should appear."""
    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove the default blank sheet

    for rows in rows_per_pdf:
        if not rows:
            continue
        customer_name = rows[0]["Customer Name"]
        ws = get_or_create_sheet(wb, customer_name)
        write_po_group(ws, rows)

    wb.save(output_path)


# ============================================================
# GUI
# ============================================================

NAVY = "#1F3864"
LIGHT_BG = "#F4F6F9"
ACCENT = "#2E7D32"


class App:
    def __init__(self, root):
        self.root = root
        root.title(APP_TITLE)
        root.geometry("620x520")
        root.minsize(560, 460)
        root.configure(bg=LIGHT_BG)

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background=LIGHT_BG)
        style.configure("TLabel", background=LIGHT_BG, font=("Segoe UI", 10))
        style.configure("Header.TLabel", background=NAVY, foreground="white",
                         font=("Segoe UI", 16, "bold"))
        style.configure("SubHeader.TLabel", background=NAVY, foreground="#D6E4F0",
                         font=("Segoe UI", 9))
        style.configure("Section.TLabel", background=LIGHT_BG, font=("Segoe UI", 10, "bold"))
        style.configure("Run.TButton", font=("Segoe UI", 11, "bold"), padding=8)
        style.configure("Secondary.TButton", font=("Segoe UI", 9), padding=4)

        self.folder_path = tk.StringVar()
        self.excel_path = tk.StringVar()
        self.status_text = tk.StringVar(value="Ready")
        self.last_output_path = None

        # ---- Header banner ----
        header = tk.Frame(root, bg=NAVY)
        header.pack(fill="x")
        tk.Label(header, text=APP_TITLE, bg=NAVY, fg="white",
                 font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=20, pady=(16, 0))
        tk.Label(header, text=APP_SUBTITLE, bg=NAVY, fg="#D6E4F0",
                 font=("Segoe UI", 9)).pack(anchor="w", padx=20, pady=(0, 14))

        # ---- Body ----
        body = ttk.Frame(root, padding=20)
        body.pack(fill="both", expand=True)

        ttk.Label(body, text="Step 1 — Select the folder containing your PO PDFs",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 6))

        folder_frame = ttk.Frame(body)
        folder_frame.pack(fill="x", pady=(0, 18))
        entry = ttk.Entry(folder_frame, textvariable=self.folder_path, font=("Segoe UI", 10))
        entry.pack(side="left", fill="x", expand=True, ipady=4)
        ttk.Button(folder_frame, text="Browse...", command=self.browse_folder).pack(
            side="left", padx=(8, 0))

        ttk.Label(body, text="Step 2 — (Optional) Use an existing Excel file",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 6))
        ttk.Label(body, text=("Leave blank to create a new po_extract.xlsx in the folder "
                               "above. Or pick an existing spreadsheet to add the new "
                               "rows into it instead."),
                  wraplength=560, foreground="#555555").pack(anchor="w", pady=(0, 6))

        excel_frame = ttk.Frame(body)
        excel_frame.pack(fill="x", pady=(0, 18))
        excel_entry = ttk.Entry(excel_frame, textvariable=self.excel_path, font=("Segoe UI", 10))
        excel_entry.pack(side="left", fill="x", expand=True, ipady=4)
        ttk.Button(excel_frame, text="Browse...", command=self.browse_excel).pack(
            side="left", padx=(8, 0))

        ttk.Label(body, text="Step 3 — Run the extraction",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 6))
        ttk.Label(body, text=("Every PDF in the folder is processed automatically. "
                               "If the output workbook already exists, new rows are "
                               "appended — each customer gets its own sheet."),
                  wraplength=560, foreground="#555555").pack(anchor="w", pady=(0, 10))

        self.run_button = ttk.Button(body, text="Run Extraction", style="Run.TButton",
                                      command=self.run_extraction)
        self.run_button.pack(fill="x", pady=(0, 14))

        ttk.Label(body, text="Activity Log", style="Section.TLabel").pack(anchor="w", pady=(0, 4))
        log_frame = tk.Frame(body, highlightbackground="#CCCCCC", highlightthickness=1)
        log_frame.pack(fill="both", expand=True)
        self.log = scrolledtext.ScrolledText(log_frame, height=10, state="disabled",
                                              font=("Consolas", 9), borderwidth=0)
        self.log.pack(fill="both", expand=True, padx=1, pady=1)

        footer = ttk.Frame(body)
        footer.pack(fill="x", pady=(12, 0))
        self.status_label = ttk.Label(footer, textvariable=self.status_text,
                                       foreground="#555555")
        self.status_label.pack(side="left")
        self.open_output_button = ttk.Button(footer, text="Open Output Folder",
                                              style="Secondary.TButton",
                                              command=self.open_output_folder,
                                              state="disabled")
        self.open_output_button.pack(side="right")

    # ---------------- UI actions ----------------

    def browse_folder(self):
        folder = filedialog.askdirectory(title="Select folder with PO PDFs")
        if folder:
            self.folder_path.set(folder)

    def browse_excel(self):
        path = filedialog.askopenfilename(
            title="Select existing Excel file",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if path:
            self.excel_path.set(path)

    def write_log(self, text):
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def open_output_folder(self):
        if not self.last_output_path:
            return
        folder = os.path.dirname(self.last_output_path)
        try:
            if sys.platform.startswith("win"):
                os.startfile(folder)
            elif sys.platform == "darwin":
                subprocess.run(["open", folder])
            else:
                subprocess.run(["xdg-open", folder])
        except Exception as e:
            messagebox.showerror("Couldn't open folder", str(e))

    def run_extraction(self):
        folder = self.folder_path.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("No folder selected", "Please choose a valid folder first.")
            return
        self.run_button.config(state="disabled")
        self.open_output_button.config(state="disabled")
        self.status_text.set("Processing...")
        threading.Thread(target=self._extract_worker, args=(folder,), daemon=True).start()

    # ---------------- Worker ----------------

    def _extract_worker(self, folder):
        pdf_files = [f for f in os.listdir(folder) if f.lower().endswith(".pdf")]
        if not pdf_files:
            self.write_log("No PDF files found in that folder.")
            self._finish("No PDFs found.")
            return

        rows_per_pdf = []
        for filename in pdf_files:
            pdf_path = os.path.join(folder, filename)
            try:
                rows = extract_fields_from_pdf(pdf_path)
                if not rows:
                    self.write_log(f"⚠  No item lines matched in: {filename}")
                rows_per_pdf.append(rows)
                self.write_log(f"✔ Processed: {filename} ({len(rows)} item rows)")
            except Exception as e:
                self.write_log(f"✘ Failed on {filename}: {e}")

        total_rows = sum(len(r) for r in rows_per_pdf)
        if total_rows == 0:
            self.write_log("Nothing extracted. No Excel file was created.")
            self._finish("Nothing extracted.")
            return

        output_path = self.excel_path.get().strip() or os.path.join(folder, "po_extract.xlsx")
        existed_before = os.path.exists(output_path)
        try:
            build_or_append_workbook(output_path, rows_per_pdf)
        except Exception as e:
            self.write_log(f"✘ Could not save Excel file: {e}")
            self._finish("Save failed.")
            return

        self.last_output_path = output_path
        verb = "Appended to" if existed_before else "Saved"
        self.write_log(f"\n{verb} {output_path}\n({total_rows} item rows across "
                        f"{len(rows_per_pdf)} PDF(s))")
        self._finish("Done.")
        self.root.after(0, lambda: self.open_output_button.config(state="normal"))

    def _finish(self, status_message):
        def apply():
            self.run_button.config(state="normal")
            self.status_text.set(status_message)
        self.root.after(0, apply)


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
