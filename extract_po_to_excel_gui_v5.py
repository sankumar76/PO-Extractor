"""
PO Extractor - desktop app version (v5 / release 1.0).

What's new in this version:
  - "About" dialog (Help menu / header button) with version and credits.
  - Successfully processed PDFs are moved into a "Processed documents"
    sub-folder inside the input folder.
  - Duplicate protection: a PO whose PO No already exists on the
    customer's sheet (or appears twice in the same run) is skipped.
  - Taller Activity Log (at least 10 visible lines).
  - Refined interface typography.
  - If the output workbook is open in Excel, it is saved and closed
    automatically before extraction starts.

Carried over from v4:
  - Output formatting (fonts, header fill color, borders, column widths)
    matches the customer-provided template exactly.
  - If the output workbook already exists, new PDFs are APPENDED to it
    rather than overwriting it (S.No numbering continues on from where
    it left off).
  - Each customer gets its own sheet/table within the workbook.

Supports TWO vendor PO formats, auto-detected per PDF:
  FORMAT A - "Accucut"-style   (labels: "PO No", "PO Date")
  FORMAT B - "Bull Agro"-style (labels: "Order No.", "Order Date")
      Order Date -> written into the "PO Date" column
      Order No.  -> written into the "PO No" column
      Description is split on the first hyphen "-"

Output columns (per customer sheet):
    S.No | PO Date | PO No | Customer Name | Item Code | Item Description |
    PO Qty | Rate | Dispatch Qty | Balance Qty
"""

import re
import os
import sys
import time
import shutil
import threading
import subprocess
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, filedialog, scrolledtext, messagebox

import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

APP_TITLE = "PO Data Extractor"
APP_SUBTITLE = "Automated Purchase Order to Excel Conversion"
APP_VERSION = "1.0"
APP_AUTHOR = "Santhosh Kumar. S"
APP_COMPANY = "ProCaliber Technologies"
APP_COPYRIGHT_YEAR = "2026"
APP_DESCRIPTION = (
    "PO Data Extractor streamlines purchase order processing by reading PO "
    "documents in PDF format and consolidating the key details - PO number, "
    "PO date, customer, item codes, descriptions, quantities and rates - into "
    "a well-structured Excel workbook.\n\n"
    "It recognises multiple vendor PO layouts automatically, organises records "
    "into a dedicated sheet per customer, guards against duplicate entries and "
    "archives every processed document, eliminating hours of manual data entry "
    "and reducing transcription errors."
)

PROCESSED_FOLDER_NAME = "Processed documents"

# ============================================================
# Extraction logic
# ============================================================

# ---------- Format A ("PO No" / "PO Date") ----------
PO_NO_PATTERN = re.compile(r"PO No\s*:\s*(\S+)")
PO_DATE_PATTERN = re.compile(r"PO Date\s*:\s*(\S+)")
ITEM_LINE_PATTERN = re.compile(
    r"^\s*\d+\s+(.+?)\s+(?:\d{6,8}\s+)?([\d.]+)\s+\w+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$"
)

# ---------- Format B ("Order No." / "Order Date") ----------
ORDER_NO_PATTERN = re.compile(r"Order No\.?\s*:\s*(.+)")
ORDER_DATE_PATTERN = re.compile(r"Order Date\s*:\s*(.+)")

LEADING_NON_ALPHA_PATTERN = re.compile(r"^([^A-Za-z]*)(.*)$", re.DOTALL)
DATE_PATTERN = re.compile(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{4})")


def normalize_date(raw_date):
    if not raw_date:
        return raw_date
    match = DATE_PATTERN.search(raw_date)
    if not match:
        return raw_date
    day, month, year = match.groups()
    return f"{int(day):02d}/{int(month):02d}/{year}"


def extract_company_name(text):
    for line in text.splitlines():
        stripped = line.strip()
        if "Limited" in stripped or "Ltd" in stripped:
            return stripped
    for line in text.splitlines():
        stripped = line.strip()
        if stripped and stripped.isupper() and not any(ch.isdigit() for ch in stripped):
            return stripped
    return ""


def split_description(description):
    """FORMAT A: split before the first alphabet character."""
    match = LEADING_NON_ALPHA_PATTERN.match(description)
    prefix = (match.group(1) or "").strip()
    remainder = (match.group(2) or "").strip()
    if prefix == "":
        return "", description.strip()
    return prefix, remainder


def split_description_hyphen(description):
    """FORMAT B: split on the first hyphen."""
    stripped = description.strip()
    if "-" in stripped:
        idx = stripped.index("-")
        return stripped[:idx].strip(), stripped[idx + 1:].strip()
    return "", stripped


def find_col_index(header_row, target):
    for i, cell in enumerate(header_row):
        if cell and target.lower() in cell.replace("\n", " ").strip().lower():
            return i
    return None


def detect_format(full_text):
    if re.search(r"PO No\s*:", full_text):
        return "A"
    if re.search(r"Order No\.?\s*:", full_text):
        return "B"
    return None


def extract_format_a(full_text):
    po_no_match = PO_NO_PATTERN.search(full_text)
    po_date_match = PO_DATE_PATTERN.search(full_text)
    po_no = po_no_match.group(1) if po_no_match else ""
    po_date = po_date_match.group(1) if po_date_match else ""

    items = []
    for line in full_text.splitlines():
        match = ITEM_LINE_PATTERN.match(line)
        if match:
            description, qty, rate, disc_amt, amount = match.groups()
            items.append((description.strip(), float(qty), float(rate)))
    return po_no, po_date, items


def extract_format_b(full_text, tables):
    order_no_match = ORDER_NO_PATTERN.search(full_text)
    order_date_match = ORDER_DATE_PATTERN.search(full_text)
    order_no = order_no_match.group(1).strip() if order_no_match else ""
    order_date = order_date_match.group(1).strip() if order_date_match else ""

    items = []
    for table in tables:
        if not table or len(table) < 3:
            continue
        header_row = table[0]
        if not header_row or not any(c and "description" in str(c).lower() for c in header_row):
            continue

        desc_idx = find_col_index(header_row, "description")
        qty_idx = find_col_index(header_row, "ord. qty") or find_col_index(header_row, "qty")
        rate_idx = find_col_index(header_row, "unit cost")
        if desc_idx is None or qty_idx is None or rate_idx is None:
            continue

        for row in table[2:]:
            sl = row[0]
            if sl and str(sl).strip().isdigit():
                raw_desc = row[desc_idx] or ""
                clean_desc = re.sub(r"HSN/SAC.*", "", raw_desc, flags=re.IGNORECASE)
                clean_desc = " ".join(clean_desc.split())
                qty_raw = (row[qty_idx] or "0").replace(",", "")
                rate_raw = (row[rate_idx] or "0").replace(",", "")
                try:
                    qty = float(qty_raw)
                    rate = float(rate_raw)
                except ValueError:
                    continue
                items.append((clean_desc, qty, rate))
    return order_no, order_date, items


def extract_fields_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        fmt = detect_format(full_text)

        if fmt == "A":
            po_no, po_date, items = extract_format_a(full_text)
            split_func = split_description
        elif fmt == "B":
            tables = []
            for page in pdf.pages:
                tables.extend(page.extract_tables())
            order_no, order_date, items = extract_format_b(full_text, tables)
            po_date, po_no = order_date, order_no
            split_func = split_description_hyphen
        else:
            raise ValueError("Unrecognized PO format (no 'PO No'/'PO Date' or "
                              "'Order No.'/'Order Date' labels found)")

    company_name = extract_company_name(full_text)
    po_date = normalize_date(po_date)

    rows = []
    for description, qty, rate in items:
        prefix, remainder = split_func(description)
        rows.append({
            "PO Date": po_date,
            "PO No": po_no,
            "Customer Name": company_name,
            "Item Code": prefix,
            "Item Description": remainder,
            "PO Qty": qty,
            "Rate": rate,
        })
    return rows


# ============================================================
# Excel output: template-matched formatting, per-customer sheets,
# append-to-existing-file support
# ============================================================

HEADERS = ["S.No", "PO Date", "PO No", "Customer Name", "Item Code",
           "Item Description", "PO Qty", "Rate", "Dispatch Qty", "Balance Qty"]

HEADER_FONT = Font(name="Arial", size=11, bold=True)
DATA_FONT = Font(name="Arial", size=11)
HEADER_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
MERGE_COLS = [1, 2, 3, 4]  # S.No, PO Date, PO No, Customer Name

COLUMN_WIDTHS = {1: 6, 2: 12.78, 3: 27.78, 4: 30.78, 5: 29.78, 6: 40.78, 7: 10.78, 9: 14.78}

INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def sanitize_sheet_name(name):
    cleaned = INVALID_SHEET_CHARS.sub("", name or "").strip()
    return (cleaned[:31] if cleaned else "Customer")


def create_sheet_with_header(wb, sheet_name):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    for col_idx, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return ws


def get_or_create_sheet(wb, customer_name):
    name = sanitize_sheet_name(customer_name)
    if name in wb.sheetnames:
        return wb[name]
    return create_sheet_with_header(wb, name)


def get_next_sno(ws):
    if ws.max_row <= 1:
        return 1
    for r in range(ws.max_row, 1, -1):
        val = ws.cell(row=r, column=1).value
        if val is not None:
            try:
                return int(val) + 1
            except (TypeError, ValueError):
                continue
    return 1


def write_po_group(ws, rows):
    """Append one PDF's rows (one PO, possibly several item rows) as a
    single group: S.No/PO Date/PO No/Customer Name centered and merged."""
    start_row = ws.max_row + 1
    s_no = get_next_sno(ws)

    for offset, row in enumerate(rows):
        r = start_row + offset
        values = [s_no, row["PO Date"], row["PO No"], row["Customer Name"],
                  row["Item Code"], row["Item Description"], row["PO Qty"],
                  row["Rate"], None, None]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx in MERGE_COLS:
                cell.alignment = CENTER_ALIGN

    end_row = start_row + len(rows) - 1
    if end_row > start_row:
        for col in MERGE_COLS:
            ws.merge_cells(start_row=start_row, start_column=col,
                            end_row=end_row, end_column=col)


def normalize_po_no(value):
    return " ".join(str(value).split()).upper() if value is not None else ""


def existing_po_numbers(ws):
    """PO numbers already on a sheet. Merged PO groups only hold the value
    in their first row, so empty cells are simply skipped."""
    found = set()
    for r in range(2, ws.max_row + 1):
        po_no = normalize_po_no(ws.cell(row=r, column=3).value)
        if po_no:
            found.add(po_no)
    return found


def build_or_append_workbook(output_path, pdf_results):
    """pdf_results: list of (filename, list_of_row_dicts) - one entry per
    successfully processed PDF, in the order they should appear.

    A PO is treated as a duplicate (and skipped) if the same PO No already
    exists on that customer's sheet, or was already added earlier in this
    run. Returns (written, duplicates) - two lists of (filename, rows)."""
    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove the default blank sheet

    known_po_numbers = {}  # sheet name -> set of normalized PO numbers
    written, duplicates = [], []

    for filename, rows in pdf_results:
        if not rows:
            continue
        sheet_name = sanitize_sheet_name(rows[0]["Customer Name"])
        if sheet_name not in known_po_numbers:
            known_po_numbers[sheet_name] = (existing_po_numbers(wb[sheet_name])
                                            if sheet_name in wb.sheetnames else set())
        seen = known_po_numbers[sheet_name]

        po_no = normalize_po_no(rows[0]["PO No"])
        if po_no and po_no in seen:
            duplicates.append((filename, rows))
            continue

        ws = get_or_create_sheet(wb, rows[0]["Customer Name"])
        write_po_group(ws, rows)
        if po_no:
            seen.add(po_no)
        written.append((filename, rows))

    if written:
        wb.save(output_path)
    return written, duplicates


# ============================================================
# File housekeeping: close the workbook in Excel, archive PDFs
# ============================================================

# Uses Excel's COM interface through Windows PowerShell so no extra Python
# packages are needed. Saves the user's pending edits, closes only the
# matching workbook, and quits Excel if nothing else is left open.
# Exit codes: 0 = workbook wasn't open, 10 = closed it, anything else = error.
_CLOSE_WORKBOOK_PS = r"""
$ErrorActionPreference = 'Stop'
$target = [IO.Path]::GetFullPath($env:PO_EXTRACTOR_TARGET)
try { $xl = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application') }
catch { exit 0 }
$closed = $false
foreach ($wb in @($xl.Workbooks)) {
    if ([string]::Equals($wb.FullName, $target, [StringComparison]::OrdinalIgnoreCase)) {
        try { $wb.Save() } catch { }
        $wb.Close($false)
        $closed = $true
    }
}
if ($closed -and $xl.Workbooks.Count -eq 0) { $xl.Quit() }
[void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
if ($closed) { exit 10 } else { exit 0 }
"""


def is_file_locked(path):
    """True if another program (e.g. Excel) holds the file open for writing."""
    if not os.path.exists(path):
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return True


def _hidden_subprocess_kwargs():
    if sys.platform.startswith("win"):
        return {"creationflags": subprocess.CREATE_NO_WINDOW}
    return {}


def _wait_until_unlocked(path, timeout=5.0):
    deadline = time.time() + timeout
    while is_file_locked(path) and time.time() < deadline:
        time.sleep(0.25)
    return not is_file_locked(path)


def close_workbook_in_excel(path):
    """Make sure `path` is not open in Excel. Returns a short message
    describing what was done, or None if nothing needed doing. Raises
    RuntimeError if the file is still locked afterwards."""
    if not sys.platform.startswith("win") or not is_file_locked(path):
        return None

    # 1) Polite close via COM: saves the user's edits first.
    env = dict(os.environ, PO_EXTRACTOR_TARGET=os.path.abspath(path))
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive",
             "-ExecutionPolicy", "Bypass", "-Command", _CLOSE_WORKBOOK_PS],
            env=env, capture_output=True, timeout=60, **_hidden_subprocess_kwargs())
        closed_via_com = result.returncode == 10
    except (OSError, subprocess.TimeoutExpired):
        closed_via_com = False

    if _wait_until_unlocked(path):
        return ("Saved and closed the open workbook in Excel."
                if closed_via_com else "Workbook was released by Excel.")

    # 2) Fallback: end the Excel window that shows this workbook.
    base = os.path.splitext(os.path.basename(path))[0]
    subprocess.run(
        ["taskkill", "/F", "/IM", "EXCEL.EXE", "/FI", f"WINDOWTITLE eq {base}*"],
        capture_output=True, **_hidden_subprocess_kwargs())
    if _wait_until_unlocked(path):
        return "Excel was force-closed to release the workbook."

    raise RuntimeError(f"'{os.path.basename(path)}' is still open in another "
                       "program. Please close it and try again.")


def move_to_processed(folder, filename):
    """Move a PDF into the 'Processed documents' sub-folder, adding a
    numeric suffix if a file with the same name is already there.
    Returns the name the file was saved under."""
    dest_dir = os.path.join(folder, PROCESSED_FOLDER_NAME)
    os.makedirs(dest_dir, exist_ok=True)
    base, ext = os.path.splitext(filename)
    dest_name = filename
    counter = 2
    while os.path.exists(os.path.join(dest_dir, dest_name)):
        dest_name = f"{base} ({counter}){ext}"
        counter += 1
    shutil.move(os.path.join(folder, filename), os.path.join(dest_dir, dest_name))
    return dest_name


# ============================================================
# GUI
# ============================================================

NAVY = "#1F3864"
LIGHT_BG = "#F4F6F9"
ACCENT = "#2E7D32"
MUTED = "#5F6B7A"
SUBTLE_TEXT = "#D6E4F0"
LOG_LINES = 14  # comfortably shows more than 10 lines without scrolling


def pick_font_family(root, candidates, fallback):
    available = set(tkfont.families(root))
    for family in candidates:
        if family in available:
            return family
    return fallback


class App:
    def __init__(self, root):
        self.root = root
        root.title(f"{APP_TITLE}  v{APP_VERSION}")
        root.geometry("760x800")
        root.minsize(640, 740)
        root.configure(bg=LIGHT_BG)

        # ---- Typography ----
        ui = pick_font_family(root, ["Segoe UI", "Helvetica Neue", "Helvetica", "Arial"],
                              "TkDefaultFont")
        ui_semibold = pick_font_family(root, ["Segoe UI Semibold"], ui)
        mono = pick_font_family(root, ["Cascadia Mono", "Consolas", "Menlo", "Courier New"],
                                "TkFixedFont")
        self.fonts = {
            "title": (ui_semibold, 18),
            "subtitle": (ui, 10),
            "section": (ui_semibold, 11),
            "body": (ui, 10),
            "hint": (ui, 9),
            "button": (ui_semibold, 11),
            "log": (mono, 10),
            "ui": ui,
            "ui_semibold": ui_semibold,
        }
        for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont"):
            try:
                tkfont.nametofont(name).configure(family=ui, size=10)
            except tk.TclError:
                pass

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background=LIGHT_BG)
        style.configure("TLabel", background=LIGHT_BG, font=self.fonts["body"])
        style.configure("Hint.TLabel", background=LIGHT_BG, foreground=MUTED,
                         font=self.fonts["hint"])
        style.configure("Section.TLabel", background=LIGHT_BG, foreground=NAVY,
                         font=self.fonts["section"])
        style.configure("TButton", font=self.fonts["body"])
        style.configure("Run.TButton", font=self.fonts["button"], padding=8)
        style.configure("Secondary.TButton", font=self.fonts["hint"], padding=4)

        self.folder_path = tk.StringVar()
        self.excel_path = tk.StringVar()
        self.status_text = tk.StringVar(value="Ready")
        self.last_output_path = None

        # ---- Menu bar ----
        menubar = tk.Menu(root)
        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label=f"About {APP_TITLE}", command=self.show_about)
        menubar.add_cascade(label="Help", menu=help_menu)
        root.config(menu=menubar)

        # ---- Header banner ----
        header = tk.Frame(root, bg=NAVY)
        header.pack(fill="x")
        header_text = tk.Frame(header, bg=NAVY)
        header_text.pack(side="left", fill="x", expand=True)
        tk.Label(header_text, text=APP_TITLE, bg=NAVY, fg="white",
                 font=self.fonts["title"]).pack(anchor="w", padx=20, pady=(14, 0))
        tk.Label(header_text, text=APP_SUBTITLE, bg=NAVY, fg=SUBTLE_TEXT,
                 font=self.fonts["subtitle"]).pack(anchor="w", padx=20, pady=(0, 14))
        about_link = tk.Label(header, text="About", bg=NAVY, fg=SUBTLE_TEXT,
                              font=(ui, 10, "underline"), cursor="hand2")
        about_link.pack(side="right", anchor="n", padx=20, pady=18)
        about_link.bind("<Button-1>", lambda _e: self.show_about())
        about_link.bind("<Enter>", lambda _e: about_link.config(fg="white"))
        about_link.bind("<Leave>", lambda _e: about_link.config(fg=SUBTLE_TEXT))

        # ---- Body ----
        body = ttk.Frame(root, padding=(20, 16, 20, 16))
        body.pack(fill="both", expand=True)

        # Footer is packed first (at the bottom) so the log never squeezes it out.
        footer = ttk.Frame(body)
        footer.pack(side="bottom", fill="x", pady=(12, 0))
        self.status_label = ttk.Label(footer, textvariable=self.status_text,
                                       style="Hint.TLabel")
        self.status_label.pack(side="left")
        self.open_output_button = ttk.Button(footer, text="Open Output Folder",
                                              style="Secondary.TButton",
                                              command=self.open_output_folder,
                                              state="disabled")
        self.open_output_button.pack(side="right")

        ttk.Label(body, text="Step 1 — Select the folder containing your PO PDFs",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 6))

        folder_frame = ttk.Frame(body)
        folder_frame.pack(fill="x", pady=(0, 16))
        entry = ttk.Entry(folder_frame, textvariable=self.folder_path, font=self.fonts["body"])
        entry.pack(side="left", fill="x", expand=True, ipady=4)
        ttk.Button(folder_frame, text="Browse...", command=self.browse_folder).pack(
            side="left", padx=(8, 0))

        ttk.Label(body, text="Step 2 — (Optional) Use an existing Excel file",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 4))
        ttk.Label(body, text=("Leave blank to create a new po_extract.xlsx in the folder "
                               "above. Or pick an existing spreadsheet to add the new "
                               "rows into it instead."),
                  wraplength=700, style="Hint.TLabel").pack(anchor="w", pady=(0, 6))

        excel_frame = ttk.Frame(body)
        excel_frame.pack(fill="x", pady=(0, 16))
        excel_entry = ttk.Entry(excel_frame, textvariable=self.excel_path, font=self.fonts["body"])
        excel_entry.pack(side="left", fill="x", expand=True, ipady=4)
        ttk.Button(excel_frame, text="Browse...", command=self.browse_excel).pack(
            side="left", padx=(8, 0))

        ttk.Label(body, text="Step 3 — Run the extraction",
                  style="Section.TLabel").pack(anchor="w", pady=(0, 4))
        ttk.Label(body, text=("Every PDF in the folder is processed automatically. New rows "
                               "are appended to the workbook (one sheet per customer), POs "
                               "already in the workbook are skipped, and processed PDFs are "
                               f"moved to a \"{PROCESSED_FOLDER_NAME}\" sub-folder."),
                  wraplength=700, style="Hint.TLabel").pack(anchor="w", pady=(0, 10))

        self.run_button = ttk.Button(body, text="Run Extraction", style="Run.TButton",
                                      command=self.run_extraction)
        self.run_button.pack(fill="x", pady=(0, 14))

        ttk.Label(body, text="Activity Log", style="Section.TLabel").pack(anchor="w", pady=(0, 4))
        log_frame = tk.Frame(body, highlightbackground="#C9CED6", highlightthickness=1)
        log_frame.pack(fill="both", expand=True)
        self.log = scrolledtext.ScrolledText(log_frame, height=LOG_LINES, state="disabled",
                                              font=self.fonts["log"], borderwidth=0,
                                              padx=6, pady=4, wrap="word")
        self.log.pack(fill="both", expand=True, padx=1, pady=1)

    # ---------------- About dialog ----------------

    def show_about(self):
        if getattr(self, "_about_win", None) and self._about_win.winfo_exists():
            self._about_win.lift()
            self._about_win.focus_force()
            return

        win = tk.Toplevel(self.root)
        self._about_win = win
        win.title(f"About {APP_TITLE}")
        win.configure(bg="white")
        win.resizable(False, False)
        win.transient(self.root)

        banner = tk.Frame(win, bg=NAVY)
        banner.pack(fill="x")
        tk.Label(banner, text=APP_TITLE, bg=NAVY, fg="white",
                 font=self.fonts["title"]).pack(anchor="w", padx=24, pady=(18, 0))
        tk.Label(banner, text=f"Version {APP_VERSION}", bg=NAVY, fg=SUBTLE_TEXT,
                 font=self.fonts["subtitle"]).pack(anchor="w", padx=24, pady=(0, 16))

        content = tk.Frame(win, bg="white")
        content.pack(fill="both", expand=True, padx=24, pady=(16, 8))

        tk.Label(content, text=APP_DESCRIPTION, bg="white", fg="#333333",
                 font=self.fonts["body"], justify="left", wraplength=440).pack(anchor="w")

        ttk.Separator(content, orient="horizontal").pack(fill="x", pady=16)

        details = tk.Frame(content, bg="white")
        details.pack(anchor="w")
        for row, (label, value) in enumerate([
            ("Version", APP_VERSION),
            ("Developed by", APP_AUTHOR),
            ("Company", APP_COMPANY),
        ]):
            tk.Label(details, text=label, bg="white", fg=MUTED,
                     font=self.fonts["body"]).grid(row=row, column=0, sticky="w",
                                                   padx=(0, 18), pady=2)
            tk.Label(details, text=value, bg="white", fg="#1A1A1A",
                     font=(self.fonts["ui_semibold"], 10)).grid(row=row, column=1,
                                                                sticky="w", pady=2)

        tk.Label(content, text=(f"© {APP_COPYRIGHT_YEAR} {APP_COMPANY}. "
                                "All rights reserved."),
                 bg="white", fg=MUTED, font=self.fonts["hint"]).pack(anchor="w", pady=(16, 0))

        button_bar = tk.Frame(win, bg="white")
        button_bar.pack(fill="x", padx=24, pady=(4, 18))
        ttk.Button(button_bar, text="OK", command=win.destroy, width=10).pack(side="right")

        win.bind("<Escape>", lambda _e: win.destroy())
        win.bind("<Return>", lambda _e: win.destroy())

        # Center over the main window.
        win.update_idletasks()
        x = self.root.winfo_rootx() + (self.root.winfo_width() - win.winfo_width()) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height() - win.winfo_height()) // 3
        win.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        win.grab_set()
        win.focus_force()

    # ---------------- UI actions ----------------

    def browse_folder(self):
        folder = filedialog.askdirectory(title="Select folder with PO PDFs")
        if folder:
            self.folder_path.set(folder)

    def browse_excel(self):
        path = filedialog.askopenfilename(
            title="Select existing Excel file",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if path:
            self.excel_path.set(path)

    def write_log(self, text):
        # Safe to call from the worker thread: the actual widget update is
        # scheduled on the Tk main loop.
        def append():
            self.log.configure(state="normal")
            self.log.insert("end", text + "\n")
            self.log.see("end")
            self.log.configure(state="disabled")
        self.root.after(0, append)

    def open_output_folder(self):
        if not self.last_output_path:
            return
        folder = os.path.dirname(self.last_output_path)
        try:
            if sys.platform.startswith("win"):
                os.startfile(folder)
            elif sys.platform == "darwin":
                subprocess.run(["open", folder])
            else:
                subprocess.run(["xdg-open", folder])
        except Exception as e:
            messagebox.showerror("Couldn't open folder", str(e))

    def run_extraction(self):
        folder = self.folder_path.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("No folder selected", "Please choose a valid folder first.")
            return
        self.run_button.config(state="disabled")
        self.open_output_button.config(state="disabled")
        self.status_text.set("Processing...")
        threading.Thread(target=self._extract_worker, args=(folder,), daemon=True).start()

    # ---------------- Worker ----------------

    def _extract_worker(self, folder):
        pdf_files = sorted(
            f for f in os.listdir(folder)
            if f.lower().endswith(".pdf") and os.path.isfile(os.path.join(folder, f))
        )
        if not pdf_files:
            self.write_log("No PDF files found in that folder.")
            self._finish("No PDFs found.")
            return

        output_path = self.excel_path.get().strip() or os.path.join(folder, "po_extract.xlsx")

        # Close the workbook in Excel first, so we can read and save it.
        try:
            message = close_workbook_in_excel(output_path)
            if message:
                self.write_log(f"ℹ {message}")
        except Exception as e:
            self.write_log(f"✘ {e}")
            self._finish("Output workbook is locked.")
            return

        pdf_results = []
        for filename in pdf_files:
            pdf_path = os.path.join(folder, filename)
            try:
                rows = extract_fields_from_pdf(pdf_path)
                if not rows:
                    self.write_log(f"⚠ No item lines matched in: {filename}")
                    continue
                pdf_results.append((filename, rows))
                self.write_log(f"✔ Processed: {filename} ({len(rows)} item rows)")
            except Exception as e:
                self.write_log(f"✘ Failed on {filename}: {e}")

        if not pdf_results:
            self.write_log("Nothing extracted. The Excel file was not changed.")
            self._finish("Nothing extracted.")
            return

        existed_before = os.path.exists(output_path)
        try:
            written, duplicates = build_or_append_workbook(output_path, pdf_results)
        except Exception as e:
            self.write_log(f"✘ Could not save Excel file: {e}")
            self._finish("Save failed.")
            return

        for filename, rows in duplicates:
            self.write_log(f"⤼ Skipped duplicate: {filename} "
                           f"(PO No {rows[0]['PO No']} already in the workbook)")

        # Archive every PDF whose data is now in the workbook (including
        # duplicates, since their data was already there).
        moved = 0
        for filename, _rows in written + duplicates:
            try:
                move_to_processed(folder, filename)
                moved += 1
            except Exception as e:
                self.write_log(f"⚠ Could not move {filename} to "
                               f"'{PROCESSED_FOLDER_NAME}': {e}")

        total_rows = sum(len(rows) for _f, rows in written)
        if written:
            self.last_output_path = output_path
            verb = "Appended to" if existed_before else "Saved"
            self.write_log(f"\n{verb} {output_path}\n({total_rows} item rows across "
                           f"{len(written)} PDF(s))")
        else:
            self.write_log("\nNo new POs to add - the Excel file was not changed.")
        if duplicates:
            self.write_log(f"{len(duplicates)} duplicate PO(s) skipped.")
        if moved:
            self.write_log(f"{moved} PDF(s) moved to '{PROCESSED_FOLDER_NAME}'.")

        self._finish("Done.")
        if written:
            self.root.after(0, lambda: self.open_output_button.config(state="normal"))

    def _finish(self, status_message):
        def apply():
            self.run_button.config(state="normal")
            self.status_text.set(status_message)
        self.root.after(0, apply)


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
