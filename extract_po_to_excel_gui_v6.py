"""
PO Extractor - desktop app version (v6 / release 1.1).

What's new in this version:
  - Default input folder:  C:/Users/<user>/Documents/POs  (created if missing)
    Default output file:   C:/Users/<user>/Documents/POs/PO_Data.xlsx
    Both can still be changed with Browse, and reset with "Use default".
  - New, bolder app logo and taskbar/window icon (see make_assets.py).
  - Redesigned, modern interface: card layout, progress bar, run summary
    (added / duplicates / failed) and a colour-coded activity log.
  - Help menu removed - it only duplicated the About dialog, which is
    still available from the header.

Carried over from v5:
  - Successfully processed PDFs are moved into a "Processed documents"
    sub-folder inside the input folder.
  - Duplicate protection: a PO whose PO No already exists on the
    customer's sheet (or appears twice in the same run) is skipped.
  - If the output workbook is open in Excel, it is saved and closed
    automatically before extraction starts.
  - Output formatting (fonts, header fill color, borders, column widths)
    matches the customer-provided template exactly.
  - If the output workbook already exists, new PDFs are APPENDED to it
    rather than overwriting it (S.No numbering continues on from where
    it left off).
  - Each customer gets its own sheet/table within the workbook.

Supports TWO vendor PO formats, auto-detected per PDF:
  FORMAT A - "Accucut"-style   (labels: "PO No", "PO Date")
  FORMAT B - "Bull Agro"-style (labels: "Order No.", "Order Date")
      Order Date -> written into the "PO Date" column
      Order No.  -> written into the "PO No" column
      Description is split on the first hyphen "-"

Output columns (per customer sheet):
    S.No | PO Date | PO No | Customer Name | Item Code | Item Description |
    PO Qty | Rate | Dispatch Qty | Balance Qty
"""

import re
import os
import sys
import time
import shutil
import threading
import subprocess
import tkinter as tk
import tkinter.font as tkfont
from datetime import datetime
from tkinter import ttk, filedialog, messagebox

import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

APP_TITLE = "PO Data Extractor"
APP_SUBTITLE = "Automated Purchase Order to Excel Conversion"
APP_VERSION = "1.1"
APP_AUTHOR = "Santhosh Kumar. S"
APP_COMPANY = "ProCaliber Technologies"
APP_COPYRIGHT_YEAR = "2026"
APP_DESCRIPTION = (
    "PO Data Extractor streamlines purchase order processing by reading PO "
    "documents in PDF format and consolidating the key details - PO number, "
    "PO date, customer, item codes, descriptions, quantities and rates - into "
    "a well-structured Excel workbook.\n\n"
    "It recognises multiple vendor PO layouts automatically, organises records "
    "into a dedicated sheet per customer, guards against duplicate entries and "
    "archives every processed document, eliminating hours of manual data entry "
    "and reducing transcription errors."
)

PROCESSED_FOLDER_NAME = "Processed documents"

DEFAULT_INPUT_FOLDER = os.path.join(os.path.expanduser("~"), "Documents", "POs")
DEFAULT_OUTPUT_FILE = os.path.join(DEFAULT_INPUT_FOLDER, "PO_Data.xlsx")

# ============================================================
# Extraction logic
# ============================================================

# ---------- Format A ("PO No" / "PO Date") ----------
PO_NO_PATTERN = re.compile(r"PO No\s*:\s*(\S+)")
PO_DATE_PATTERN = re.compile(r"PO Date\s*:\s*(\S+)")
ITEM_LINE_PATTERN = re.compile(
    r"^\s*\d+\s+(.+?)\s+(?:\d{6,8}\s+)?([\d.]+)\s+\w+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$"
)

# ---------- Format B ("Order No." / "Order Date") ----------
ORDER_NO_PATTERN = re.compile(r"Order No\.?\s*:\s*(.+)")
ORDER_DATE_PATTERN = re.compile(r"Order Date\s*:\s*(.+)")

LEADING_NON_ALPHA_PATTERN = re.compile(r"^([^A-Za-z]*)(.*)$", re.DOTALL)
DATE_PATTERN = re.compile(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{4})")


def normalize_date(raw_date):
    if not raw_date:
        return raw_date
    match = DATE_PATTERN.search(raw_date)
    if not match:
        return raw_date
    day, month, year = match.groups()
    return f"{int(day):02d}/{int(month):02d}/{year}"


def extract_company_name(text):
    for line in text.splitlines():
        stripped = line.strip()
        if "Limited" in stripped or "Ltd" in stripped:
            return stripped
    for line in text.splitlines():
        stripped = line.strip()
        if stripped and stripped.isupper() and not any(ch.isdigit() for ch in stripped):
            return stripped
    return ""


def split_description(description):
    """FORMAT A: split before the first alphabet character."""
    match = LEADING_NON_ALPHA_PATTERN.match(description)
    prefix = (match.group(1) or "").strip()
    remainder = (match.group(2) or "").strip()
    if prefix == "":
        return "", description.strip()
    return prefix, remainder


def split_description_hyphen(description):
    """FORMAT B: split on the first hyphen."""
    stripped = description.strip()
    if "-" in stripped:
        idx = stripped.index("-")
        return stripped[:idx].strip(), stripped[idx + 1:].strip()
    return "", stripped


def find_col_index(header_row, target):
    for i, cell in enumerate(header_row):
        if cell and target.lower() in cell.replace("\n", " ").strip().lower():
            return i
    return None


def detect_format(full_text):
    if re.search(r"PO No\s*:", full_text):
        return "A"
    if re.search(r"Order No\.?\s*:", full_text):
        return "B"
    return None


def extract_format_a(full_text):
    po_no_match = PO_NO_PATTERN.search(full_text)
    po_date_match = PO_DATE_PATTERN.search(full_text)
    po_no = po_no_match.group(1) if po_no_match else ""
    po_date = po_date_match.group(1) if po_date_match else ""

    items = []
    for line in full_text.splitlines():
        match = ITEM_LINE_PATTERN.match(line)
        if match:
            description, qty, rate, disc_amt, amount = match.groups()
            items.append((description.strip(), float(qty), float(rate)))
    return po_no, po_date, items


def extract_format_b(full_text, tables):
    order_no_match = ORDER_NO_PATTERN.search(full_text)
    order_date_match = ORDER_DATE_PATTERN.search(full_text)
    order_no = order_no_match.group(1).strip() if order_no_match else ""
    order_date = order_date_match.group(1).strip() if order_date_match else ""

    items = []
    for table in tables:
        if not table or len(table) < 3:
            continue
        header_row = table[0]
        if not header_row or not any(c and "description" in str(c).lower() for c in header_row):
            continue

        desc_idx = find_col_index(header_row, "description")
        qty_idx = find_col_index(header_row, "ord. qty") or find_col_index(header_row, "qty")
        rate_idx = find_col_index(header_row, "unit cost")
        if desc_idx is None or qty_idx is None or rate_idx is None:
            continue

        for row in table[2:]:
            sl = row[0]
            if sl and str(sl).strip().isdigit():
                raw_desc = row[desc_idx] or ""
                clean_desc = re.sub(r"HSN/SAC.*", "", raw_desc, flags=re.IGNORECASE)
                clean_desc = " ".join(clean_desc.split())
                qty_raw = (row[qty_idx] or "0").replace(",", "")
                rate_raw = (row[rate_idx] or "0").replace(",", "")
                try:
                    qty = float(qty_raw)
                    rate = float(rate_raw)
                except ValueError:
                    continue
                items.append((clean_desc, qty, rate))
    return order_no, order_date, items


def extract_fields_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        fmt = detect_format(full_text)

        if fmt == "A":
            po_no, po_date, items = extract_format_a(full_text)
            split_func = split_description
        elif fmt == "B":
            tables = []
            for page in pdf.pages:
                tables.extend(page.extract_tables())
            order_no, order_date, items = extract_format_b(full_text, tables)
            po_date, po_no = order_date, order_no
            split_func = split_description_hyphen
        else:
            raise ValueError("Unrecognized PO format (no 'PO No'/'PO Date' or "
                              "'Order No.'/'Order Date' labels found)")

    company_name = extract_company_name(full_text)
    po_date = normalize_date(po_date)

    rows = []
    for description, qty, rate in items:
        prefix, remainder = split_func(description)
        rows.append({
            "PO Date": po_date,
            "PO No": po_no,
            "Customer Name": company_name,
            "Item Code": prefix,
            "Item Description": remainder,
            "PO Qty": qty,
            "Rate": rate,
        })
    return rows


# ============================================================
# Excel output: template-matched formatting, per-customer sheets,
# append-to-existing-file support
# ============================================================

HEADERS = ["S.No", "PO Date", "PO No", "Customer Name", "Item Code",
           "Item Description", "PO Qty", "Rate", "Dispatch Qty", "Balance Qty"]

HEADER_FONT = Font(name="Arial", size=11, bold=True)
DATA_FONT = Font(name="Arial", size=11)
HEADER_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
MERGE_COLS = [1, 2, 3, 4]  # S.No, PO Date, PO No, Customer Name

COLUMN_WIDTHS = {1: 6, 2: 12.78, 3: 27.78, 4: 30.78, 5: 29.78, 6: 40.78, 7: 10.78, 9: 14.78}

INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def sanitize_sheet_name(name):
    cleaned = INVALID_SHEET_CHARS.sub("", name or "").strip()
    return (cleaned[:31] if cleaned else "Customer")


def create_sheet_with_header(wb, sheet_name):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    for col_idx, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return ws


def get_or_create_sheet(wb, customer_name):
    name = sanitize_sheet_name(customer_name)
    if name in wb.sheetnames:
        return wb[name]
    return create_sheet_with_header(wb, name)


def get_next_sno(ws):
    if ws.max_row <= 1:
        return 1
    for r in range(ws.max_row, 1, -1):
        val = ws.cell(row=r, column=1).value
        if val is not None:
            try:
                return int(val) + 1
            except (TypeError, ValueError):
                continue
    return 1


def write_po_group(ws, rows):
    """Append one PDF's rows (one PO, possibly several item rows) as a
    single group: S.No/PO Date/PO No/Customer Name centered and merged."""
    start_row = ws.max_row + 1
    s_no = get_next_sno(ws)

    for offset, row in enumerate(rows):
        r = start_row + offset
        values = [s_no, row["PO Date"], row["PO No"], row["Customer Name"],
                  row["Item Code"], row["Item Description"], row["PO Qty"],
                  row["Rate"], None, None]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx in MERGE_COLS:
                cell.alignment = CENTER_ALIGN

    end_row = start_row + len(rows) - 1
    if end_row > start_row:
        for col in MERGE_COLS:
            ws.merge_cells(start_row=start_row, start_column=col,
                            end_row=end_row, end_column=col)


def normalize_po_no(value):
    return " ".join(str(value).split()).upper() if value is not None else ""


def existing_po_numbers(ws):
    """PO numbers already on a sheet. Merged PO groups only hold the value
    in their first row, so empty cells are simply skipped."""
    found = set()
    for r in range(2, ws.max_row + 1):
        po_no = normalize_po_no(ws.cell(row=r, column=3).value)
        if po_no:
            found.add(po_no)
    return found


def build_or_append_workbook(output_path, pdf_results):
    """pdf_results: list of (filename, list_of_row_dicts) - one entry per
    successfully processed PDF, in the order they should appear.

    A PO is treated as a duplicate (and skipped) if the same PO No already
    exists on that customer's sheet, or was already added earlier in this
    run. Returns (written, duplicates) - two lists of (filename, rows)."""
    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove the default blank sheet

    known_po_numbers = {}  # sheet name -> set of normalized PO numbers
    written, duplicates = [], []

    for filename, rows in pdf_results:
        if not rows:
            continue
        sheet_name = sanitize_sheet_name(rows[0]["Customer Name"])
        if sheet_name not in known_po_numbers:
            known_po_numbers[sheet_name] = (existing_po_numbers(wb[sheet_name])
                                            if sheet_name in wb.sheetnames else set())
        seen = known_po_numbers[sheet_name]

        po_no = normalize_po_no(rows[0]["PO No"])
        if po_no and po_no in seen:
            duplicates.append((filename, rows))
            continue

        ws = get_or_create_sheet(wb, rows[0]["Customer Name"])
        write_po_group(ws, rows)
        if po_no:
            seen.add(po_no)
        written.append((filename, rows))

    if written:
        wb.save(output_path)
    return written, duplicates


# ============================================================
# File housekeeping: close the workbook in Excel, archive PDFs
# ============================================================

# Uses Excel's COM interface through Windows PowerShell so no extra Python
# packages are needed. Saves the user's pending edits, closes only the
# matching workbook, and quits Excel if nothing else is left open.
# Exit codes: 0 = workbook wasn't open, 10 = closed it, anything else = error.
_CLOSE_WORKBOOK_PS = r"""
$ErrorActionPreference = 'Stop'
$target = [IO.Path]::GetFullPath($env:PO_EXTRACTOR_TARGET)
try { $xl = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application') }
catch { exit 0 }
$closed = $false
foreach ($wb in @($xl.Workbooks)) {
    if ([string]::Equals($wb.FullName, $target, [StringComparison]::OrdinalIgnoreCase)) {
        try { $wb.Save() } catch { }
        $wb.Close($false)
        $closed = $true
    }
}
if ($closed -and $xl.Workbooks.Count -eq 0) { $xl.Quit() }
[void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
if ($closed) { exit 10 } else { exit 0 }
"""


def is_file_locked(path):
    """True if another program (e.g. Excel) holds the file open for writing."""
    if not os.path.exists(path):
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return True


def _hidden_subprocess_kwargs():
    if sys.platform.startswith("win"):
        return {"creationflags": subprocess.CREATE_NO_WINDOW}
    return {}


def _wait_until_unlocked(path, timeout=5.0):
    deadline = time.time() + timeout
    while is_file_locked(path) and time.time() < deadline:
        time.sleep(0.25)
    return not is_file_locked(path)


def close_workbook_in_excel(path):
    """Make sure `path` is not open in Excel. Returns a short message
    describing what was done, or None if nothing needed doing. Raises
    RuntimeError if the file is still locked afterwards."""
    if not sys.platform.startswith("win") or not is_file_locked(path):
        return None

    # 1) Polite close via COM: saves the user's edits first.
    env = dict(os.environ, PO_EXTRACTOR_TARGET=os.path.abspath(path))
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive",
             "-ExecutionPolicy", "Bypass", "-Command", _CLOSE_WORKBOOK_PS],
            env=env, capture_output=True, timeout=60, **_hidden_subprocess_kwargs())
        closed_via_com = result.returncode == 10
    except (OSError, subprocess.TimeoutExpired):
        closed_via_com = False

    if _wait_until_unlocked(path):
        return ("Saved and closed the open workbook in Excel."
                if closed_via_com else "Workbook was released by Excel.")

    # 2) Fallback: end the Excel window that shows this workbook.
    base = os.path.splitext(os.path.basename(path))[0]
    subprocess.run(
        ["taskkill", "/F", "/IM", "EXCEL.EXE", "/FI", f"WINDOWTITLE eq {base}*"],
        capture_output=True, **_hidden_subprocess_kwargs())
    if _wait_until_unlocked(path):
        return "Excel was force-closed to release the workbook."

    raise RuntimeError(f"'{os.path.basename(path)}' is still open in another "
                       "program. Please close it and try again.")


def move_to_processed(folder, filename):
    """Move a PDF into the 'Processed documents' sub-folder, adding a
    numeric suffix if a file with the same name is already there.
    Returns the name the file was saved under."""
    dest_dir = os.path.join(folder, PROCESSED_FOLDER_NAME)
    os.makedirs(dest_dir, exist_ok=True)
    base, ext = os.path.splitext(filename)
    dest_name = filename
    counter = 2
    while os.path.exists(os.path.join(dest_dir, dest_name)):
        dest_name = f"{base} ({counter}){ext}"
        counter += 1
    shutil.move(os.path.join(folder, filename), os.path.join(dest_dir, dest_name))
    return dest_name


# ============================================================
# GUI
# ============================================================

# Palette
BG = "#EEF2F7"
CARD = "#FFFFFF"
BORDER = "#DCE3EC"
NAVY = "#1E3A8A"
PRIMARY = "#2563EB"
PRIMARY_HOVER = "#1D4ED8"
PRIMARY_DISABLED = "#93B4F5"
TEXT = "#0F172A"
MUTED = "#64748B"
SUBTLE_TEXT = "#DBEAFE"
SECONDARY_BG = "#E8EEF8"
SECONDARY_HOVER = "#D6E1F5"
SUCCESS = "#16A34A"
WARNING = "#D97706"
ERROR = "#DC2626"
LOG_BG = "#0F172A"
LOG_FG = "#E2E8F0"
LOG_MUTED = "#64748B"
LOG_LINES = 12  # comfortably shows more than 10 lines without scrolling

APP_USER_MODEL_ID = "ProCaliber.PODataExtractor"
LOGO_SIZES = (48, 56, 64, 72, 80, 96, 112, 128)  # generated by make_assets.py


def resource_path(*parts):
    """Path to a bundled file - works both from source and from the
    PyInstaller one-file .exe (which unpacks to sys._MEIPASS)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, *parts)


def prepare_windows_process():
    """Crisp text on high-DPI screens, and our own taskbar icon/grouping
    instead of the generic Python one. Must run before Tk is created."""
    if not sys.platform.startswith("win"):
        return
    import ctypes
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass
    try:
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(APP_USER_MODEL_ID)
    except Exception:
        pass


def pick_font_family(root, candidates, fallback):
    available = set(tkfont.families(root))
    for family in candidates:
        if family in available:
            return family
    return fallback


class FlatButton(tk.Label):
    """A flat, modern button with hover and disabled states."""

    def __init__(self, parent, text, command, bg, fg, hover_bg, font,
                 disabled_bg=None, disabled_fg=None, padx=16, pady=8):
        super().__init__(parent, text=text, bg=bg, fg=fg, font=font,
                         padx=padx, pady=pady, cursor="hand2")
        self._command = command
        self._colors = (bg, fg, hover_bg, disabled_bg or bg, disabled_fg or MUTED)
        self._enabled = True
        self.bind("<Enter>", lambda _e: self._enabled and self.config(bg=self._colors[2]))
        self.bind("<Leave>", lambda _e: self._enabled and self.config(bg=self._colors[0]))
        self.bind("<ButtonRelease-1>", self._on_click)

    def _on_click(self, event):
        inside = 0 <= event.x <= self.winfo_width() and 0 <= event.y <= self.winfo_height()
        if self._enabled and inside and self._command:
            self._command()

    def set_enabled(self, enabled):
        self._enabled = enabled
        bg, fg, _hover, dis_bg, dis_fg = self._colors
        self.config(bg=bg if enabled else dis_bg, fg=fg if enabled else dis_fg,
                    cursor="hand2" if enabled else "arrow")


class App:
    def __init__(self, root):
        self.root = root
        self.scale = max(1.0, root.winfo_fpixels("1i") / 96.0)
        root.title(f"{APP_TITLE}  v{APP_VERSION}")
        # Fit on smaller / high-DPI screens (leave room for the taskbar).
        usable_height = root.winfo_screenheight() - self.px(90)
        width, height = self.px(840), min(self.px(860), usable_height)
        x = max((root.winfo_screenwidth() - width) // 2, 0)
        root.geometry(f"{width}x{height}+{x}+{self.px(10)}")
        root.minsize(self.px(700), min(self.px(720), usable_height))
        root.configure(bg=BG)
        self._set_window_icon()

        # ---- Typography ----
        ui = pick_font_family(root, ["Segoe UI", "Helvetica Neue", "Helvetica", "Arial"],
                              "TkDefaultFont")
        ui_semibold = pick_font_family(root, ["Segoe UI Semibold"], ui)
        mono = pick_font_family(root, ["Cascadia Mono", "Consolas", "Menlo", "Courier New"],
                                "TkFixedFont")
        self.fonts = {
            "title": (ui_semibold, 19),
            "subtitle": (ui, 10),
            "section": (ui_semibold, 12),
            "body": (ui, 10),
            "hint": (ui, 9),
            "button": (ui_semibold, 10),
            "run": (ui_semibold, 12),
            "stat_value": (ui_semibold, 16),
            "badge": (ui_semibold, 10),
            "log": (mono, 9),
            "ui": ui,
            "ui_semibold": ui_semibold,
        }
        for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont"):
            try:
                tkfont.nametofont(name).configure(family=ui, size=10)
            except tk.TclError:
                pass

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Modern.Horizontal.TProgressbar", troughcolor=SECONDARY_BG,
                        background=PRIMARY, bordercolor=SECONDARY_BG,
                        lightcolor=PRIMARY, darkcolor=PRIMARY, thickness=self.px(6))
        style.configure("Log.Vertical.TScrollbar", troughcolor=LOG_BG, background="#334155",
                        bordercolor=LOG_BG, arrowcolor=LOG_FG, lightcolor="#334155",
                        darkcolor="#334155", gripcount=0)

        self.folder_path = tk.StringVar(value=DEFAULT_INPUT_FOLDER)
        self.excel_path = tk.StringVar(value=DEFAULT_OUTPUT_FILE)
        self.status_text = tk.StringVar(value="Ready")
        self.last_output_path = None
        self._running = False

        # Make sure the default input folder exists so users can drop PDFs in.
        try:
            os.makedirs(DEFAULT_INPUT_FOLDER, exist_ok=True)
        except OSError:
            pass

        self._build_header()

        # Status bar is packed first (at the bottom) so the log never squeezes it out.
        self._build_status_bar()

        body = tk.Frame(root, bg=BG)
        body.pack(fill="both", expand=True, padx=self.px(20), pady=(self.px(12), 0))

        self._build_path_card(
            body, 1, "Input folder",
            "The folder containing your PO PDFs. Processed PDFs are moved to a "
            f"\"{PROCESSED_FOLDER_NAME}\" sub-folder inside it.",
            self.folder_path, self.browse_folder, DEFAULT_INPUT_FOLDER)
        self._build_path_card(
            body, 2, "Output Excel file",
            "New POs are appended to this workbook (one sheet per customer). "
            "It is created if it doesn't exist yet; POs already in it are skipped.",
            self.excel_path, self.browse_excel, DEFAULT_OUTPUT_FILE)
        self._build_run_card(body)
        self._build_log_card(body)

    # ---------------- Helpers ----------------

    def px(self, value):
        return int(round(value * self.scale))

    def _set_window_icon(self):
        ico = resource_path("assets", "app.ico")
        try:
            if sys.platform.startswith("win") and os.path.exists(ico):
                self.root.iconbitmap(default=ico)  # applies to every window
            else:
                self.root.iconphoto(True, tk.PhotoImage(file=resource_path("assets",
                                                                           "logo_96.png")))
        except tk.TclError:
            pass

    def _load_logo(self, base_size):
        """Load the logo PNG that best matches base_size at the current DPI."""
        target = base_size * self.scale
        size = min(LOGO_SIZES, key=lambda s: abs(s - target))
        try:
            return tk.PhotoImage(file=resource_path("assets", f"logo_{size}.png"))
        except tk.TclError:
            return None

    def _card(self, parent, expand=False):
        card = tk.Frame(parent, bg=CARD, highlightbackground=BORDER, highlightthickness=1)
        card.pack(fill="both" if expand else "x", expand=expand, pady=(0, self.px(10)))
        inner = tk.Frame(card, bg=CARD)
        inner.pack(fill="both", expand=True, padx=self.px(18), pady=self.px(11))
        return inner

    def _step_badge(self, parent, number):
        size = self.px(26)
        c = tk.Canvas(parent, width=size, height=size, bg=CARD, highlightthickness=0)
        c.create_oval(1, 1, size - 1, size - 1, fill=PRIMARY, outline=PRIMARY)
        c.create_text(size / 2, size / 2, text=str(number), fill="white",
                      font=self.fonts["badge"])
        return c

    def _link(self, parent, text, command, bg=CARD, fg=PRIMARY, hover=PRIMARY_HOVER):
        link = tk.Label(parent, text=text, bg=bg, fg=fg, cursor="hand2",
                        font=(self.fonts["ui"], 9, "underline"))
        link.bind("<Button-1>", lambda _e: command())
        link.bind("<Enter>", lambda _e: link.config(fg=hover))
        link.bind("<Leave>", lambda _e: link.config(fg=fg))
        return link

    def _secondary_button(self, parent, text, command):
        return FlatButton(parent, text, command, bg=SECONDARY_BG, fg=NAVY,
                          hover_bg=SECONDARY_HOVER, font=self.fonts["button"],
                          disabled_fg="#A0AEC0", padx=self.px(14), pady=self.px(6))

    # ---------------- Layout ----------------

    def _build_header(self):
        height = self.px(78)
        header = tk.Canvas(self.root, height=height, highlightthickness=0, bg=NAVY)
        header.pack(fill="x")

        self.header_logo = self._load_logo(56)
        x = self.px(22)
        if self.header_logo:
            header.create_image(x, height // 2, image=self.header_logo, anchor="w")
            x += self.header_logo.width() + self.px(14)
        header.create_text(x, height // 2 + self.px(2), text=APP_TITLE, anchor="sw",
                           fill="white", font=self.fonts["title"])
        header.create_text(x, height // 2 + self.px(4), text=APP_SUBTITLE, anchor="nw",
                           fill=SUBTLE_TEXT, font=self.fonts["subtitle"])

        about = FlatButton(header, "About", self.show_about, bg="#2D4FB3", fg="white",
                           hover_bg="#3B5FCC", font=self.fonts["button"],
                           padx=self.px(14), pady=self.px(5))
        about_id = header.create_window(0, height // 2, window=about, anchor="e")

        def redraw(event):
            header.delete("grad")
            w = event.width
            left, right = (30, 58, 138), (37, 99, 235)  # navy -> blue, left to right
            step = 4
            for i in range(0, w, step):
                t = i / max(w - 1, 1)
                color = "#%02x%02x%02x" % tuple(int(left[k] + (right[k] - left[k]) * t)
                                                for k in range(3))
                header.create_rectangle(i, 0, i + step, height, fill=color, width=0,
                                        tags="grad")
            header.tag_lower("grad")
            header.coords(about_id, w - self.px(22), height // 2)

        header.bind("<Configure>", redraw)

    def _build_path_card(self, parent, number, title, hint, variable, browse, default):
        inner = self._card(parent)

        title_row = tk.Frame(inner, bg=CARD)
        title_row.pack(fill="x")
        self._step_badge(title_row, number).pack(side="left")
        tk.Label(title_row, text=title, bg=CARD, fg=TEXT,
                 font=self.fonts["section"]).pack(side="left", padx=(self.px(10), 0))
        self._link(title_row, "Use default", lambda: variable.set(default)).pack(side="right")

        tk.Label(inner, text=hint, bg=CARD, fg=MUTED, font=self.fonts["hint"],
                 justify="left", anchor="w", wraplength=self.px(740)).pack(
            fill="x", pady=(self.px(3), self.px(6)))

        row = tk.Frame(inner, bg=CARD)
        row.pack(fill="x")
        entry_wrap = tk.Frame(row, bg="#F8FAFC", highlightbackground=BORDER,
                              highlightcolor=PRIMARY, highlightthickness=1)
        entry_wrap.pack(side="left", fill="x", expand=True)
        entry = tk.Entry(entry_wrap, textvariable=variable, relief="flat", bg="#F8FAFC",
                         fg=TEXT, insertbackground=TEXT, font=self.fonts["body"],
                         highlightthickness=0, borderwidth=0)
        entry.pack(fill="x", ipady=self.px(5), padx=(self.px(8), self.px(4)))
        entry.bind("<FocusIn>", lambda _e: entry_wrap.config(highlightbackground=PRIMARY))
        entry.bind("<FocusOut>", lambda _e: entry_wrap.config(highlightbackground=BORDER))
        self._secondary_button(row, "Browse…", browse).pack(side="left",
                                                            padx=(self.px(10), 0))

    def _build_run_card(self, parent):
        inner = self._card(parent)

        top = tk.Frame(inner, bg=CARD)
        top.pack(fill="x")
        self.run_button = FlatButton(top, "▶  Run Extraction", self.run_extraction,
                                     bg=PRIMARY, fg="white", hover_bg=PRIMARY_HOVER,
                                     disabled_bg=PRIMARY_DISABLED, disabled_fg="white",
                                     font=self.fonts["run"],
                                     padx=self.px(26), pady=self.px(8))
        self.run_button.pack(side="left")

        stats = tk.Frame(top, bg=CARD)
        stats.pack(side="right")
        self.stat_vars = {}
        for key, label, color in (("added", "POs added", SUCCESS),
                                  ("duplicates", "Duplicates", WARNING),
                                  ("failed", "Failed", ERROR)):
            box = tk.Frame(stats, bg=CARD)
            box.pack(side="left", padx=(self.px(22), 0))
            var = tk.StringVar(value="–")
            self.stat_vars[key] = var
            tk.Label(box, textvariable=var, bg=CARD, fg=color,
                     font=self.fonts["stat_value"]).pack(anchor="e")
            tk.Label(box, text=label, bg=CARD, fg=MUTED,
                     font=self.fonts["hint"]).pack(anchor="e")

        self.progress = ttk.Progressbar(inner, style="Modern.Horizontal.TProgressbar",
                                        mode="determinate", maximum=1, value=0)
        self.progress.pack(fill="x", pady=(self.px(10), 0))

    def _build_log_card(self, parent):
        inner = self._card(parent, expand=True)

        title_row = tk.Frame(inner, bg=CARD)
        title_row.pack(fill="x", pady=(0, self.px(6)))
        tk.Label(title_row, text="Activity Log", bg=CARD, fg=TEXT,
                 font=self.fonts["section"]).pack(side="left")
        self._link(title_row, "Clear", self.clear_log).pack(side="right")

        log_frame = tk.Frame(inner, bg=LOG_BG)
        log_frame.pack(fill="both", expand=True)
        scrollbar = ttk.Scrollbar(log_frame, orient="vertical",
                                  style="Log.Vertical.TScrollbar")
        scrollbar.pack(side="right", fill="y")
        self.log = tk.Text(log_frame, height=LOG_LINES, state="disabled", bg=LOG_BG,
                           fg=LOG_FG, insertbackground=LOG_FG, font=self.fonts["log"],
                           borderwidth=0, highlightthickness=0, wrap="word",
                           padx=self.px(12), pady=self.px(8),
                           yscrollcommand=scrollbar.set)
        self.log.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.log.yview)
        self.log.tag_configure("time", foreground=LOG_MUTED)
        self.log.tag_configure("ok", foreground="#4ADE80")
        self.log.tag_configure("warn", foreground="#FBBF24")
        self.log.tag_configure("err", foreground="#F87171")
        self.log.tag_configure("info", foreground="#93C5FD")

        self.write_log(f"Welcome to {APP_TITLE}. Check the input folder and output "
                       "file above, then click Run Extraction.")

    def _build_status_bar(self):
        bar = tk.Frame(self.root, bg=CARD, highlightbackground=BORDER, highlightthickness=1)
        bar.pack(side="bottom", fill="x")
        inner = tk.Frame(bar, bg=CARD)
        inner.pack(fill="x", padx=self.px(20), pady=self.px(6))

        self.status_dot = tk.Label(inner, text="●", bg=CARD, fg=SUCCESS,
                                   font=self.fonts["body"])
        self.status_dot.pack(side="left")
        tk.Label(inner, textvariable=self.status_text, bg=CARD, fg=MUTED,
                 font=self.fonts["hint"]).pack(side="left", padx=(self.px(6), 0))

        self.open_excel_button = self._secondary_button(inner, "Open Excel File",
                                                        self.open_output_file)
        self.open_excel_button.pack(side="right")
        self.open_output_button = self._secondary_button(inner, "Open Output Folder",
                                                         self.open_output_folder)
        self.open_output_button.pack(side="right", padx=(0, self.px(8)))
        self._set_output_buttons()

    # ---------------- About dialog ----------------

    def show_about(self):
        if getattr(self, "_about_win", None) and self._about_win.winfo_exists():
            self._about_win.lift()
            self._about_win.focus_force()
            return

        win = tk.Toplevel(self.root)
        self._about_win = win
        win.title(f"About {APP_TITLE}")
        win.configure(bg=CARD)
        win.resizable(False, False)
        win.transient(self.root)

        banner = tk.Frame(win, bg=NAVY)
        banner.pack(fill="x")
        self.about_logo = self._load_logo(72)
        if self.about_logo:
            tk.Label(banner, image=self.about_logo, bg=NAVY).pack(
                side="left", padx=(self.px(24), self.px(16)), pady=self.px(18))
        text = tk.Frame(banner, bg=NAVY)
        text.pack(side="left", pady=self.px(18), padx=(0, self.px(24)))
        tk.Label(text, text=APP_TITLE, bg=NAVY, fg="white",
                 font=self.fonts["title"]).pack(anchor="w")
        tk.Label(text, text=f"Version {APP_VERSION}", bg=NAVY, fg=SUBTLE_TEXT,
                 font=self.fonts["subtitle"]).pack(anchor="w")

        content = tk.Frame(win, bg=CARD)
        content.pack(fill="both", expand=True, padx=self.px(24),
                     pady=(self.px(16), self.px(8)))

        tk.Label(content, text=APP_DESCRIPTION, bg=CARD, fg="#334155",
                 font=self.fonts["body"], justify="left",
                 wraplength=self.px(460)).pack(anchor="w")

        tk.Frame(content, bg=BORDER, height=1).pack(fill="x", pady=self.px(16))

        details = tk.Frame(content, bg=CARD)
        details.pack(anchor="w")
        for row, (label, value) in enumerate([
            ("Version", APP_VERSION),
            ("Developed by", APP_AUTHOR),
            ("Company", APP_COMPANY),
        ]):
            tk.Label(details, text=label, bg=CARD, fg=MUTED,
                     font=self.fonts["body"]).grid(row=row, column=0, sticky="w",
                                                   padx=(0, self.px(18)), pady=2)
            tk.Label(details, text=value, bg=CARD, fg=TEXT,
                     font=(self.fonts["ui_semibold"], 10)).grid(row=row, column=1,
                                                                sticky="w", pady=2)

        tk.Label(content, text=(f"© {APP_COPYRIGHT_YEAR} {APP_COMPANY}. "
                                "All rights reserved."),
                 bg=CARD, fg=MUTED, font=self.fonts["hint"]).pack(anchor="w",
                                                                  pady=(self.px(16), 0))

        button_bar = tk.Frame(win, bg=CARD)
        button_bar.pack(fill="x", padx=self.px(24), pady=(self.px(4), self.px(18)))
        FlatButton(button_bar, "OK", win.destroy, bg=PRIMARY, fg="white",
                   hover_bg=PRIMARY_HOVER, font=self.fonts["button"],
                   padx=self.px(24), pady=self.px(6)).pack(side="right")

        win.bind("<Escape>", lambda _e: win.destroy())
        win.bind("<Return>", lambda _e: win.destroy())

        # Center over the main window.
        win.update_idletasks()
        x = self.root.winfo_rootx() + (self.root.winfo_width() - win.winfo_width()) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height() - win.winfo_height()) // 3
        win.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        win.grab_set()
        win.focus_force()

    # ---------------- UI actions ----------------

    def browse_folder(self):
        current = self.folder_path.get().strip()
        folder = filedialog.askdirectory(
            title="Select folder with PO PDFs",
            initialdir=current if os.path.isdir(current) else os.path.expanduser("~"))
        if folder:
            self.folder_path.set(os.path.normpath(folder))

    def browse_excel(self):
        current = self.excel_path.get().strip() or DEFAULT_OUTPUT_FILE
        initial_dir = os.path.dirname(current)
        path = filedialog.asksaveasfilename(
            title="Choose output Excel file (new or existing)",
            initialdir=initial_dir if os.path.isdir(initial_dir) else os.path.expanduser("~"),
            initialfile=os.path.basename(current),
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            confirmoverwrite=False,  # existing files are appended to, never overwritten
        )
        if path:
            self.excel_path.set(os.path.normpath(path))

    def write_log(self, text):
        # Safe to call from the worker thread: the actual widget update is
        # scheduled on the Tk main loop. Lines are coloured by their leading
        # symbol (✔ ok, ⚠/⤼ warning, ✘ error, ℹ info).
        stripped = text.lstrip()
        lead = text[:len(text) - len(stripped)]  # keeps leading blank lines
        tag = {"✔": "ok", "⚠": "warn", "⤼": "warn", "✘": "err", "ℹ": "info"}.get(
            stripped[:1], ())
        stamp = datetime.now().strftime("%H:%M:%S")

        def append():
            self.log.configure(state="normal")
            self.log.insert("end", lead)
            self.log.insert("end", f"{stamp}  ", "time")
            self.log.insert("end", stripped + "\n", tag)
            self.log.see("end")
            self.log.configure(state="disabled")
        self.root.after(0, append)

    def clear_log(self):
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

    def _set_output_buttons(self):
        available = bool(self.last_output_path) and not self._running
        self.open_output_button.set_enabled(available)
        self.open_excel_button.set_enabled(available
                                           and os.path.exists(self.last_output_path))

    def _open_path(self, path):
        try:
            if sys.platform.startswith("win"):
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.run(["open", path])
            else:
                subprocess.run(["xdg-open", path])
        except Exception as e:
            messagebox.showerror("Couldn't open", str(e))

    def open_output_folder(self):
        if self.last_output_path:
            self._open_path(os.path.dirname(self.last_output_path))

    def open_output_file(self):
        if self.last_output_path and os.path.exists(self.last_output_path):
            self._open_path(self.last_output_path)

    def _set_status(self, message, color):
        self.status_text.set(message)
        self.status_dot.config(fg=color)

    def run_extraction(self):
        folder = self.folder_path.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("Input folder not found",
                                 "Please choose a valid input folder first.")
            return
        output_path = self.excel_path.get().strip() or DEFAULT_OUTPUT_FILE
        if not output_path.lower().endswith(".xlsx"):
            output_path += ".xlsx"
        output_path = os.path.abspath(output_path)
        self.excel_path.set(output_path)
        try:
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
        except OSError as e:
            messagebox.showerror("Output folder not available", str(e))
            return

        self._running = True
        self.run_button.set_enabled(False)
        self.run_button.config(text="Processing…")
        self._set_output_buttons()
        for var in self.stat_vars.values():
            var.set("–")
        self.progress.config(value=0, maximum=1)
        self._set_status("Processing…", WARNING)
        threading.Thread(target=self._extract_worker, args=(folder, output_path),
                         daemon=True).start()

    # ---------------- Worker ----------------

    def _progress(self, done, total):
        self.root.after(0, lambda: self.progress.config(maximum=max(total, 1), value=done))

    def _extract_worker(self, folder, output_path):
        pdf_files = sorted(
            f for f in os.listdir(folder)
            if f.lower().endswith(".pdf") and os.path.isfile(os.path.join(folder, f))
        )
        if not pdf_files:
            self.write_log(f"⚠ No PDF files found in {folder}")
            self._finish("No PDFs found.", WARNING)
            return

        self.write_log(f"Found {len(pdf_files)} PDF(s) in {folder}")

        # Close the workbook in Excel first, so we can read and save it.
        try:
            message = close_workbook_in_excel(output_path)
            if message:
                self.write_log(f"ℹ {message}")
        except Exception as e:
            self.write_log(f"✘ {e}")
            self._finish("Output workbook is locked.", ERROR)
            return

        pdf_results = []
        failed = 0
        for index, filename in enumerate(pdf_files, start=1):
            pdf_path = os.path.join(folder, filename)
            try:
                rows = extract_fields_from_pdf(pdf_path)
                if rows:
                    pdf_results.append((filename, rows))
                    self.write_log(f"✔ Processed: {filename} ({len(rows)} item rows)")
                else:
                    failed += 1
                    self.write_log(f"⚠ No item lines matched in: {filename}")
            except Exception as e:
                failed += 1
                self.write_log(f"✘ Failed on {filename}: {e}")
            self._progress(index, len(pdf_files))

        if not pdf_results:
            self.write_log("Nothing extracted. The Excel file was not changed.")
            self._finish("Nothing extracted.", ERROR, (0, 0, failed))
            return

        existed_before = os.path.exists(output_path)
        try:
            written, duplicates = build_or_append_workbook(output_path, pdf_results)
        except Exception as e:
            self.write_log(f"✘ Could not save Excel file: {e}")
            self._finish("Save failed.", ERROR, (0, 0, failed))
            return

        for filename, rows in duplicates:
            self.write_log(f"⤼ Skipped duplicate: {filename} "
                           f"(PO No {rows[0]['PO No']} already in the workbook)")

        # Archive every PDF whose data is now in the workbook (including
        # duplicates, since their data was already there).
        moved = 0
        for filename, _rows in written + duplicates:
            try:
                move_to_processed(folder, filename)
                moved += 1
            except Exception as e:
                self.write_log(f"⚠ Could not move {filename} to "
                               f"'{PROCESSED_FOLDER_NAME}': {e}")

        total_rows = sum(len(rows) for _f, rows in written)
        if written:
            self.last_output_path = output_path
            verb = "Appended to" if existed_before else "Saved"
            self.write_log(f"\n✔ {verb} {output_path} ({total_rows} item rows across "
                           f"{len(written)} PDF(s))")
        else:
            self.write_log("\nℹ No new POs to add - the Excel file was not changed.")
        if duplicates:
            self.write_log(f"{len(duplicates)} duplicate PO(s) skipped.")
        if moved:
            self.write_log(f"{moved} PDF(s) moved to '{PROCESSED_FOLDER_NAME}'.")

        if failed:
            self._finish(f"Done with {failed} problem file(s).", WARNING,
                         (len(written), len(duplicates), failed))
        else:
            self._finish("Done.", SUCCESS, (len(written), len(duplicates), failed))

    def _finish(self, status_message, color=SUCCESS, counts=None):
        def apply():
            self._running = False
            self.run_button.set_enabled(True)
            self.run_button.config(text="▶  Run Extraction")
            self._set_status(status_message, color)
            if counts is not None:
                for key, value in zip(("added", "duplicates", "failed"), counts):
                    self.stat_vars[key].set(str(value))
            self._set_output_buttons()
        self.root.after(0, apply)


if __name__ == "__main__":
    prepare_windows_process()
    root = tk.Tk()
    app = App(root)
    root.mainloop()
