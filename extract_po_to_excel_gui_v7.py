"""
PO Extractor - desktop app version (v7 / release 1.1).

What's new in this version:
  - Professional interface rebuilt on Qt (PySide6): neutral light theme
    with a single green accent, rounded controls, two-column layout,
    run summary tiles and an activity table.
  - Default input folder:  C:/Users/<user>/Documents/POs  (created if missing)
    Default output file:   C:/Users/<user>/Documents/POs/PO_Data.xlsx
    Both can still be changed with Browse, or put back with "Reset to default".
  - App logo and taskbar/window icon (see make_assets.py).
  - Help menu removed - it only duplicated the About dialog, which is
    available from the header.

Carried over from v5:
  - Successfully processed PDFs are moved into a "Processed documents"
    sub-folder inside the input folder.
  - Duplicate protection: a PO whose PO No already exists on the
    customer's sheet (or appears twice in the same run) is skipped.
  - If the output workbook is open in Excel, it is saved and closed
    automatically before extraction starts.
  - Output formatting (fonts, header fill color, borders, column widths)
    matches the customer-provided template exactly.
  - If the output workbook already exists, new PDFs are APPENDED to it
    rather than overwriting it (S.No numbering continues on from where
    it left off).
  - Each customer gets its own sheet/table within the workbook.

Supports TWO vendor PO formats, auto-detected per PDF:
  FORMAT A - "Accucut"-style   (labels: "PO No", "PO Date")
  FORMAT B - "Bull Agro"-style (labels: "Order No.", "Order Date")
      Order Date -> written into the "PO Date" column
      Order No.  -> written into the "PO No" column
      Description is split on the first hyphen "-"

Output columns (per customer sheet):
    S.No | PO Date | PO No | Customer Name | Item Code | Item Description |
    PO Qty | Rate | Dispatch Qty | Balance Qty
"""

import re
import os
import sys
import time
import shutil
import threading
import subprocess
from datetime import datetime

import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt, QObject, Signal
from PySide6.QtGui import QColor, QFont, QIcon, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QDialog, QFileDialog, QFrame,
    QGraphicsDropShadowEffect, QGridLayout, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QMainWindow, QMessageBox, QProgressBar, QPushButton, QStatusBar,
    QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

APP_TITLE = "PO Data Extractor"
APP_SUBTITLE = "Automated Purchase Order to Excel Conversion"
APP_VERSION = "1.1"
APP_AUTHOR = "Santhosh Kumar. S"
APP_COMPANY = "ProCaliber Technologies"
APP_COPYRIGHT_YEAR = "2026"
APP_DESCRIPTION = (
    "PO Data Extractor streamlines purchase order processing by reading PO "
    "documents in PDF format and consolidating the key details - PO number, "
    "PO date, customer, item codes, descriptions, quantities and rates - into "
    "a well-structured Excel workbook.\n\n"
    "It recognises multiple vendor PO layouts automatically, organises records "
    "into a dedicated sheet per customer, guards against duplicate entries and "
    "archives every processed document, eliminating hours of manual data entry "
    "and reducing transcription errors."
)

PROCESSED_FOLDER_NAME = "Processed documents"

DEFAULT_INPUT_FOLDER = os.path.join(os.path.expanduser("~"), "Documents", "POs")
DEFAULT_OUTPUT_FILE = os.path.join(DEFAULT_INPUT_FOLDER, "PO_Data.xlsx")

# ============================================================
# Extraction logic
# ============================================================

# ---------- Format A ("PO No" / "PO Date") ----------
PO_NO_PATTERN = re.compile(r"PO No\s*:\s*(\S+)")
PO_DATE_PATTERN = re.compile(r"PO Date\s*:\s*(\S+)")
ITEM_LINE_PATTERN = re.compile(
    r"^\s*\d+\s+(.+?)\s+(?:\d{6,8}\s+)?([\d.]+)\s+\w+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$"
)

# ---------- Format B ("Order No." / "Order Date") ----------
ORDER_NO_PATTERN = re.compile(r"Order No\.?\s*:\s*(.+)")
ORDER_DATE_PATTERN = re.compile(r"Order Date\s*:\s*(.+)")

LEADING_NON_ALPHA_PATTERN = re.compile(r"^([^A-Za-z]*)(.*)$", re.DOTALL)
DATE_PATTERN = re.compile(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{4})")


def normalize_date(raw_date):
    if not raw_date:
        return raw_date
    match = DATE_PATTERN.search(raw_date)
    if not match:
        return raw_date
    day, month, year = match.groups()
    return f"{int(day):02d}/{int(month):02d}/{year}"


def extract_company_name(text):
    for line in text.splitlines():
        stripped = line.strip()
        if "Limited" in stripped or "Ltd" in stripped:
            return stripped
    for line in text.splitlines():
        stripped = line.strip()
        if stripped and stripped.isupper() and not any(ch.isdigit() for ch in stripped):
            return stripped
    return ""


def split_description(description):
    """FORMAT A: split before the first alphabet character."""
    match = LEADING_NON_ALPHA_PATTERN.match(description)
    prefix = (match.group(1) or "").strip()
    remainder = (match.group(2) or "").strip()
    if prefix == "":
        return "", description.strip()
    return prefix, remainder


def split_description_hyphen(description):
    """FORMAT B: split on the first hyphen."""
    stripped = description.strip()
    if "-" in stripped:
        idx = stripped.index("-")
        return stripped[:idx].strip(), stripped[idx + 1:].strip()
    return "", stripped


def find_col_index(header_row, target):
    for i, cell in enumerate(header_row):
        if cell and target.lower() in cell.replace("\n", " ").strip().lower():
            return i
    return None


def detect_format(full_text):
    if re.search(r"PO No\s*:", full_text):
        return "A"
    if re.search(r"Order No\.?\s*:", full_text):
        return "B"
    return None


def extract_format_a(full_text):
    po_no_match = PO_NO_PATTERN.search(full_text)
    po_date_match = PO_DATE_PATTERN.search(full_text)
    po_no = po_no_match.group(1) if po_no_match else ""
    po_date = po_date_match.group(1) if po_date_match else ""

    items = []
    for line in full_text.splitlines():
        match = ITEM_LINE_PATTERN.match(line)
        if match:
            description, qty, rate, disc_amt, amount = match.groups()
            items.append((description.strip(), float(qty), float(rate)))
    return po_no, po_date, items


def extract_format_b(full_text, tables):
    order_no_match = ORDER_NO_PATTERN.search(full_text)
    order_date_match = ORDER_DATE_PATTERN.search(full_text)
    order_no = order_no_match.group(1).strip() if order_no_match else ""
    order_date = order_date_match.group(1).strip() if order_date_match else ""

    items = []
    for table in tables:
        if not table or len(table) < 3:
            continue
        header_row = table[0]
        if not header_row or not any(c and "description" in str(c).lower() for c in header_row):
            continue

        desc_idx = find_col_index(header_row, "description")
        qty_idx = find_col_index(header_row, "ord. qty") or find_col_index(header_row, "qty")
        rate_idx = find_col_index(header_row, "unit cost")
        if desc_idx is None or qty_idx is None or rate_idx is None:
            continue

        for row in table[2:]:
            sl = row[0]
            if sl and str(sl).strip().isdigit():
                raw_desc = row[desc_idx] or ""
                clean_desc = re.sub(r"HSN/SAC.*", "", raw_desc, flags=re.IGNORECASE)
                clean_desc = " ".join(clean_desc.split())
                qty_raw = (row[qty_idx] or "0").replace(",", "")
                rate_raw = (row[rate_idx] or "0").replace(",", "")
                try:
                    qty = float(qty_raw)
                    rate = float(rate_raw)
                except ValueError:
                    continue
                items.append((clean_desc, qty, rate))
    return order_no, order_date, items


def extract_fields_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        fmt = detect_format(full_text)

        if fmt == "A":
            po_no, po_date, items = extract_format_a(full_text)
            split_func = split_description
        elif fmt == "B":
            tables = []
            for page in pdf.pages:
                tables.extend(page.extract_tables())
            order_no, order_date, items = extract_format_b(full_text, tables)
            po_date, po_no = order_date, order_no
            split_func = split_description_hyphen
        else:
            raise ValueError("Unrecognized PO format (no 'PO No'/'PO Date' or "
                              "'Order No.'/'Order Date' labels found)")

    company_name = extract_company_name(full_text)
    po_date = normalize_date(po_date)

    rows = []
    for description, qty, rate in items:
        prefix, remainder = split_func(description)
        rows.append({
            "PO Date": po_date,
            "PO No": po_no,
            "Customer Name": company_name,
            "Item Code": prefix,
            "Item Description": remainder,
            "PO Qty": qty,
            "Rate": rate,
        })
    return rows


# ============================================================
# Excel output: template-matched formatting, per-customer sheets,
# append-to-existing-file support
# ============================================================

HEADERS = ["S.No", "PO Date", "PO No", "Customer Name", "Item Code",
           "Item Description", "PO Qty", "Rate", "Dispatch Qty", "Balance Qty"]

HEADER_FONT = Font(name="Arial", size=11, bold=True)
DATA_FONT = Font(name="Arial", size=11)
HEADER_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
MERGE_COLS = [1, 2, 3, 4]  # S.No, PO Date, PO No, Customer Name

COLUMN_WIDTHS = {1: 6, 2: 12.78, 3: 27.78, 4: 30.78, 5: 29.78, 6: 40.78, 7: 10.78, 9: 14.78}

INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def sanitize_sheet_name(name):
    cleaned = INVALID_SHEET_CHARS.sub("", name or "").strip()
    return (cleaned[:31] if cleaned else "Customer")


def create_sheet_with_header(wb, sheet_name):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    for col_idx, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return ws


def get_or_create_sheet(wb, customer_name):
    name = sanitize_sheet_name(customer_name)
    if name in wb.sheetnames:
        return wb[name]
    return create_sheet_with_header(wb, name)


def get_next_sno(ws):
    if ws.max_row <= 1:
        return 1
    for r in range(ws.max_row, 1, -1):
        val = ws.cell(row=r, column=1).value
        if val is not None:
            try:
                return int(val) + 1
            except (TypeError, ValueError):
                continue
    return 1


def write_po_group(ws, rows):
    """Append one PDF's rows (one PO, possibly several item rows) as a
    single group: S.No/PO Date/PO No/Customer Name centered and merged."""
    start_row = ws.max_row + 1
    s_no = get_next_sno(ws)

    for offset, row in enumerate(rows):
        r = start_row + offset
        values = [s_no, row["PO Date"], row["PO No"], row["Customer Name"],
                  row["Item Code"], row["Item Description"], row["PO Qty"],
                  row["Rate"], None, None]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx in MERGE_COLS:
                cell.alignment = CENTER_ALIGN

    end_row = start_row + len(rows) - 1
    if end_row > start_row:
        for col in MERGE_COLS:
            ws.merge_cells(start_row=start_row, start_column=col,
                            end_row=end_row, end_column=col)


def normalize_po_no(value):
    return " ".join(str(value).split()).upper() if value is not None else ""


def existing_po_numbers(ws):
    """PO numbers already on a sheet. Merged PO groups only hold the value
    in their first row, so empty cells are simply skipped."""
    found = set()
    for r in range(2, ws.max_row + 1):
        po_no = normalize_po_no(ws.cell(row=r, column=3).value)
        if po_no:
            found.add(po_no)
    return found


def build_or_append_workbook(output_path, pdf_results):
    """pdf_results: list of (filename, list_of_row_dicts) - one entry per
    successfully processed PDF, in the order they should appear.

    A PO is treated as a duplicate (and skipped) if the same PO No already
    exists on that customer's sheet, or was already added earlier in this
    run. Returns (written, duplicates) - two lists of (filename, rows)."""
    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove the default blank sheet

    known_po_numbers = {}  # sheet name -> set of normalized PO numbers
    written, duplicates = [], []

    for filename, rows in pdf_results:
        if not rows:
            continue
        sheet_name = sanitize_sheet_name(rows[0]["Customer Name"])
        if sheet_name not in known_po_numbers:
            known_po_numbers[sheet_name] = (existing_po_numbers(wb[sheet_name])
                                            if sheet_name in wb.sheetnames else set())
        seen = known_po_numbers[sheet_name]

        po_no = normalize_po_no(rows[0]["PO No"])
        if po_no and po_no in seen:
            duplicates.append((filename, rows))
            continue

        ws = get_or_create_sheet(wb, rows[0]["Customer Name"])
        write_po_group(ws, rows)
        if po_no:
            seen.add(po_no)
        written.append((filename, rows))

    if written:
        wb.save(output_path)
    return written, duplicates


# ============================================================
# File housekeeping: close the workbook in Excel, archive PDFs
# ============================================================

# Uses Excel's COM interface through Windows PowerShell so no extra Python
# packages are needed. Saves the user's pending edits, closes only the
# matching workbook, and quits Excel if nothing else is left open.
# Exit codes: 0 = workbook wasn't open, 10 = closed it, anything else = error.
_CLOSE_WORKBOOK_PS = r"""
$ErrorActionPreference = 'Stop'
$target = [IO.Path]::GetFullPath($env:PO_EXTRACTOR_TARGET)
try { $xl = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application') }
catch { exit 0 }
$closed = $false
foreach ($wb in @($xl.Workbooks)) {
    if ([string]::Equals($wb.FullName, $target, [StringComparison]::OrdinalIgnoreCase)) {
        try { $wb.Save() } catch { }
        $wb.Close($false)
        $closed = $true
    }
}
if ($closed -and $xl.Workbooks.Count -eq 0) { $xl.Quit() }
[void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
if ($closed) { exit 10 } else { exit 0 }
"""


def is_file_locked(path):
    """True if another program (e.g. Excel) holds the file open for writing."""
    if not os.path.exists(path):
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return True


def _hidden_subprocess_kwargs():
    if sys.platform.startswith("win"):
        return {"creationflags": subprocess.CREATE_NO_WINDOW}
    return {}


def _wait_until_unlocked(path, timeout=5.0):
    deadline = time.time() + timeout
    while is_file_locked(path) and time.time() < deadline:
        time.sleep(0.25)
    return not is_file_locked(path)


def close_workbook_in_excel(path):
    """Make sure `path` is not open in Excel. Returns a short message
    describing what was done, or None if nothing needed doing. Raises
    RuntimeError if the file is still locked afterwards."""
    if not sys.platform.startswith("win") or not is_file_locked(path):
        return None

    # 1) Polite close via COM: saves the user's edits first.
    env = dict(os.environ, PO_EXTRACTOR_TARGET=os.path.abspath(path))
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive",
             "-ExecutionPolicy", "Bypass", "-Command", _CLOSE_WORKBOOK_PS],
            env=env, capture_output=True, timeout=60, **_hidden_subprocess_kwargs())
        closed_via_com = result.returncode == 10
    except (OSError, subprocess.TimeoutExpired):
        closed_via_com = False

    if _wait_until_unlocked(path):
        return ("Saved and closed the open workbook in Excel."
                if closed_via_com else "Workbook was released by Excel.")

    # 2) Fallback: end the Excel window that shows this workbook.
    base = os.path.splitext(os.path.basename(path))[0]
    subprocess.run(
        ["taskkill", "/F", "/IM", "EXCEL.EXE", "/FI", f"WINDOWTITLE eq {base}*"],
        capture_output=True, **_hidden_subprocess_kwargs())
    if _wait_until_unlocked(path):
        return "Excel was force-closed to release the workbook."

    raise RuntimeError(f"'{os.path.basename(path)}' is still open in another "
                       "program. Please close it and try again.")


def move_to_processed(folder, filename):
    """Move a PDF into the 'Processed documents' sub-folder, adding a
    numeric suffix if a file with the same name is already there.
    Returns the name the file was saved under."""
    dest_dir = os.path.join(folder, PROCESSED_FOLDER_NAME)
    os.makedirs(dest_dir, exist_ok=True)
    base, ext = os.path.splitext(filename)
    dest_name = filename
    counter = 2
    while os.path.exists(os.path.join(dest_dir, dest_name)):
        dest_name = f"{base} ({counter}){ext}"
        counter += 1
    shutil.move(os.path.join(folder, filename), os.path.join(dest_dir, dest_name))
    return dest_name


# ============================================================
# GUI (Qt / PySide6)
# ============================================================

# Neutral, enterprise palette with a single green accent.
WINDOW_BG = "#F3F4F6"
SURFACE = "#FFFFFF"
SURFACE_ALT = "#FAFBFC"
BORDER = "#E2E5E9"
INPUT_BORDER = "#C9CED6"
INPUT_BORDER_HOVER = "#9AA3AE"
TEXT = "#1F2328"
TEXT_SECONDARY = "#5B6470"
TEXT_DISABLED = "#A0A7B0"
ACCENT = "#107C41"
ACCENT_HOVER = "#0C6534"
ACCENT_PRESSED = "#0A5129"
ACCENT_DISABLED = "#8FC2A6"
ACCENT_SOFT = "#E6F2EB"

LEVEL_STYLES = {  # activity table: level -> (label, colour)
    "success": ("Success", "#107C41"),
    "skipped": ("Skipped", "#B26B00"),
    "warning": ("Warning", "#B26B00"),
    "error": ("Error", "#C42B1C"),
    "info": ("Info", "#3B6EA8"),
}

APP_USER_MODEL_ID = "ProCaliber.PODataExtractor"

STYLESHEET = f"""
QWidget {{
    color: {TEXT};
    font-family: "Segoe UI";
    font-size: 9pt;
}}
QMainWindow, #central {{ background: {WINDOW_BG}; }}

#header {{ background: {SURFACE}; border-bottom: 1px solid {BORDER}; }}
#appTitle {{ font-size: 14pt; font-weight: 600; }}
#appSubtitle {{ color: {TEXT_SECONDARY}; }}

#card {{ background: {SURFACE}; border: 1px solid {BORDER}; border-radius: 8px; }}
#cardTitle {{ font-size: 11pt; font-weight: 600; }}
#muted {{ color: {TEXT_SECONDARY}; }}
#fieldLabel {{ font-weight: 600; }}

QLineEdit {{
    background: {SURFACE};
    border: 1px solid {INPUT_BORDER};
    border-radius: 6px;
    padding: 6px 10px;
    selection-background-color: {ACCENT};
}}
QLineEdit:hover {{ border-color: {INPUT_BORDER_HOVER}; }}
QLineEdit:focus {{ border: 1px solid {ACCENT}; }}

QPushButton {{
    background: {SURFACE};
    border: 1px solid {INPUT_BORDER};
    border-radius: 6px;
    padding: 6px 14px;
    font-weight: 600;
}}
QPushButton:hover {{ background: {WINDOW_BG}; border-color: {INPUT_BORDER_HOVER}; }}
QPushButton:pressed {{ background: #E7E9EC; }}
QPushButton:disabled {{ color: {TEXT_DISABLED}; background: {SURFACE_ALT}; border-color: {BORDER}; }}

QPushButton#primary {{
    background: {ACCENT};
    border: 1px solid {ACCENT};
    color: white;
    padding: 9px 22px;
    font-size: 10pt;
}}
QPushButton#primary:hover {{ background: {ACCENT_HOVER}; border-color: {ACCENT_HOVER}; }}
QPushButton#primary:pressed {{ background: {ACCENT_PRESSED}; }}
QPushButton#primary:disabled {{ background: {ACCENT_DISABLED}; border-color: {ACCENT_DISABLED}; }}

QPushButton#link {{
    background: transparent;
    border: none;
    color: {ACCENT};
    padding: 2px 2px;
}}
QPushButton#link:hover {{ color: {ACCENT_HOVER}; text-decoration: underline; }}

QProgressBar {{
    background: #E7E9EC;
    border: none;
    border-radius: 3px;
    max-height: 6px;
    min-height: 6px;
}}
QProgressBar::chunk {{ background: {ACCENT}; border-radius: 3px; }}

#kpi {{ background: {SURFACE_ALT}; border: 1px solid {BORDER}; border-radius: 6px; }}
#kpiLabel {{ color: {TEXT_SECONDARY}; font-size: 8pt; font-weight: 600; }}
QLabel#kpiValue {{ font-size: 17pt; font-weight: 600; color: {TEXT}; }}
QLabel#kpiValue[tone="empty"] {{ color: {TEXT_DISABLED}; }}
QLabel#kpiValue[tone="success"] {{ color: {LEVEL_STYLES["success"][1]}; }}
QLabel#kpiValue[tone="skipped"] {{ color: {LEVEL_STYLES["skipped"][1]}; }}
QLabel#kpiValue[tone="error"] {{ color: {LEVEL_STYLES["error"][1]}; }}

QTableWidget {{
    background: {SURFACE};
    alternate-background-color: {SURFACE_ALT};
    border: 1px solid {BORDER};
    border-radius: 6px;
    gridline-color: transparent;
    selection-background-color: {ACCENT_SOFT};
    selection-color: {TEXT};
}}
QTableWidget::item {{ padding-left: 8px; padding-right: 8px; }}
QHeaderView::section {{
    background: {SURFACE_ALT};
    color: {TEXT_SECONDARY};
    border: none;
    border-bottom: 1px solid {BORDER};
    padding: 6px 8px;
    font-weight: 600;
}}

QScrollBar:vertical {{ background: transparent; width: 10px; margin: 2px; }}
QScrollBar::handle:vertical {{ background: {INPUT_BORDER}; border-radius: 3px; min-height: 28px; }}
QScrollBar::handle:vertical:hover {{ background: {INPUT_BORDER_HOVER}; }}
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0; }}
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{ background: none; }}

QStatusBar {{ background: {SURFACE}; border-top: 1px solid {BORDER}; color: {TEXT_SECONDARY}; }}
QStatusBar::item {{ border: none; }}

QToolTip {{ background: {TEXT}; color: white; border: none; padding: 4px 8px; }}

QDialog {{ background: {SURFACE}; }}
"""


def resource_path(*parts):
    """Path to a bundled file - works both from source and from the
    PyInstaller one-file .exe (which unpacks to sys._MEIPASS)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, *parts)


def set_app_user_model_id():
    """Give the app its own taskbar icon/grouping instead of Python's."""
    if sys.platform.startswith("win"):
        import ctypes
        try:
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(APP_USER_MODEL_ID)
        except Exception:
            pass


def logo_pixmap(size, device_pixel_ratio):
    """The app logo at `size` logical pixels, sharp on high-DPI screens."""
    pixmap = QPixmap(resource_path("assets", "logo_128.png"))
    if pixmap.isNull():
        return pixmap
    physical = int(size * device_pixel_ratio)
    pixmap = pixmap.scaled(physical, physical, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    pixmap.setDevicePixelRatio(device_pixel_ratio)
    return pixmap


def make_card():
    card = QFrame()
    card.setObjectName("card")
    shadow = QGraphicsDropShadowEffect(card)
    shadow.setBlurRadius(18)
    shadow.setOffset(0, 2)
    shadow.setColor(QColor(16, 24, 40, 18))
    card.setGraphicsEffect(shadow)
    layout = QVBoxLayout(card)
    layout.setContentsMargins(20, 18, 20, 20)
    layout.setSpacing(12)
    return card, layout


def make_label(text, object_name=None, wrap=False):
    label = QLabel(text)
    if object_name:
        label.setObjectName(object_name)
    label.setWordWrap(wrap)
    return label


def make_button(text, object_name=None, slot=None):
    button = QPushButton(text)
    if object_name:
        button.setObjectName(object_name)
    button.setCursor(Qt.PointingHandCursor)
    if slot:
        button.clicked.connect(slot)
    return button


class WorkerSignals(QObject):
    """Bridges the background extraction thread to the UI thread."""
    log = Signal(str, str)          # level, message
    progress = Signal(int, int)     # done, total
    finished = Signal(dict)         # summary


class KpiTile(QFrame):
    def __init__(self, label, tone):
        super().__init__()
        self.setObjectName("kpi")
        self._tone = tone  # colour used when the value is non-zero
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 8, 12, 8)
        layout.setSpacing(0)
        layout.addWidget(make_label(label.upper(), "kpiLabel"))
        self.value = make_label("–", "kpiValue")
        layout.addWidget(self.value)
        self.set_value(None)

    def set_value(self, value):
        self.value.setText("–" if value is None else str(value))
        tone = "empty" if value is None else (self._tone if value else "")
        self.value.setProperty("tone", tone)
        self.value.style().unpolish(self.value)
        self.value.style().polish(self.value)


class AboutDialog(QDialog):
    def __init__(self, parent):
        super().__init__(parent)
        self.setWindowTitle(f"About {APP_TITLE}")
        self.setFixedWidth(520)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(28, 26, 28, 22)
        layout.setSpacing(16)

        top = QHBoxLayout()
        top.setSpacing(16)
        logo = QLabel()
        logo.setPixmap(logo_pixmap(64, self.devicePixelRatioF()))
        top.addWidget(logo, 0, Qt.AlignTop)
        titles = QVBoxLayout()
        titles.setSpacing(2)
        titles.addWidget(make_label(APP_TITLE, "appTitle"))
        titles.addWidget(make_label(f"Version {APP_VERSION}", "muted"))
        titles.addStretch()
        top.addLayout(titles, 1)
        layout.addLayout(top)

        description = make_label(APP_DESCRIPTION, wrap=True)
        description.setStyleSheet("line-height: 140%;")
        layout.addWidget(description)

        divider = QFrame()
        divider.setFixedHeight(1)
        divider.setStyleSheet(f"background: {BORDER};")
        layout.addWidget(divider)

        details = QGridLayout()
        details.setHorizontalSpacing(24)
        details.setVerticalSpacing(6)
        for row, (label, value) in enumerate([("Version", APP_VERSION),
                                              ("Developed by", APP_AUTHOR),
                                              ("Company", APP_COMPANY)]):
            details.addWidget(make_label(label, "muted"), row, 0)
            details.addWidget(make_label(value, "fieldLabel"), row, 1)
        details.setColumnStretch(1, 1)
        layout.addLayout(details)

        layout.addWidget(make_label(f"© {APP_COPYRIGHT_YEAR} {APP_COMPANY}. "
                                    "All rights reserved.", "muted"))

        buttons = QHBoxLayout()
        buttons.addStretch()
        ok = make_button("OK", "primary", self.accept)
        ok.setDefault(True)
        ok.setMinimumWidth(96)
        buttons.addWidget(ok)
        layout.addLayout(buttons)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_TITLE}  v{APP_VERSION}")
        self.last_output_path = None
        self._running = False
        self._about = None

        self.signals = WorkerSignals()
        self.signals.log.connect(self.add_log)
        self.signals.progress.connect(self.on_progress)
        self.signals.finished.connect(self.on_finished)

        # Make sure the default input folder exists so users can drop PDFs in.
        try:
            os.makedirs(DEFAULT_INPUT_FOLDER, exist_ok=True)
        except OSError:
            pass

        central = QWidget()
        central.setObjectName("central")
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        root.addWidget(self._build_header())

        body = QHBoxLayout()
        body.setContentsMargins(20, 20, 20, 16)
        body.setSpacing(16)
        left = QVBoxLayout()
        left.setSpacing(16)
        left.addWidget(self._build_files_card())
        left.addWidget(self._build_run_card())
        left.addStretch()
        left_widget = QWidget()
        left_widget.setLayout(left)
        left_widget.setFixedWidth(430)
        left.setContentsMargins(0, 0, 0, 0)
        body.addWidget(left_widget)
        body.addWidget(self._build_activity_card(), 1)
        root.addLayout(body, 1)

        self._build_status_bar()
        self.add_log("info", f"Ready. Place PO PDFs in the input folder and click "
                             f"Run Extraction.")

        # Size to fit the screen (leaves room for the taskbar).
        screen = QApplication.primaryScreen().availableGeometry()
        width, height = min(1120, screen.width() - 40), min(760, screen.height() - 40)
        self.resize(width, height)
        self.setMinimumSize(min(900, width), min(620, height))
        self.move(screen.x() + (screen.width() - width) // 2,
                  screen.y() + (screen.height() - height) // 2)

    # ---------------- Layout ----------------

    def _build_header(self):
        header = QFrame()
        header.setObjectName("header")
        layout = QHBoxLayout(header)
        layout.setContentsMargins(24, 14, 24, 14)
        layout.setSpacing(14)

        logo = QLabel()
        logo.setPixmap(logo_pixmap(44, self.devicePixelRatioF()))
        layout.addWidget(logo)

        titles = QVBoxLayout()
        titles.setSpacing(0)
        titles.addWidget(make_label(APP_TITLE, "appTitle"))
        titles.addWidget(make_label(APP_SUBTITLE, "appSubtitle"))
        layout.addLayout(titles, 1)

        layout.addWidget(make_button("About", slot=self.show_about))
        return header

    def _path_field(self, label, hint, default, browse_slot):
        block = QVBoxLayout()
        block.setSpacing(6)

        top = QHBoxLayout()
        top.addWidget(make_label(label, "fieldLabel"))
        top.addStretch()
        top.addWidget(make_button("Reset to default", "link", lambda: edit.setText(default)))
        block.addLayout(top)

        row = QHBoxLayout()
        row.setSpacing(8)
        edit = QLineEdit(default)
        edit.setClearButtonEnabled(True)
        edit.setMinimumHeight(34)
        row.addWidget(edit, 1)
        browse = make_button("Browse…", slot=browse_slot)
        browse.setMinimumHeight(34)
        row.addWidget(browse)
        block.addLayout(row)

        block.addWidget(make_label(hint, "muted", wrap=True))
        return block, edit

    def _build_files_card(self):
        card, layout = make_card()
        layout.addWidget(make_label("Source & destination", "cardTitle"))

        block, self.folder_edit = self._path_field(
            "Input folder",
            f"PDFs are moved to \"{PROCESSED_FOLDER_NAME}\" after processing.",
            DEFAULT_INPUT_FOLDER, self.browse_folder)
        layout.addLayout(block)
        layout.addSpacing(4)
        block, self.excel_edit = self._path_field(
            "Output workbook",
            "Created if missing. New POs are appended; one sheet per customer.",
            DEFAULT_OUTPUT_FILE, self.browse_excel)
        layout.addLayout(block)
        return card

    def _build_run_card(self):
        card, layout = make_card()
        layout.addWidget(make_label("Extraction", "cardTitle"))
        layout.addWidget(make_label("Processes every PDF in the input folder. POs already "
                                    "in the workbook are skipped automatically.",
                                    "muted", wrap=True))

        self.run_button = make_button("Run Extraction", "primary", self.run_extraction)
        self.run_button.setMinimumHeight(40)
        layout.addWidget(self.run_button)

        self.progress = QProgressBar()
        self.progress.setTextVisible(False)
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        layout.addWidget(self.progress)
        self.progress_label = make_label("Ready", "muted")
        layout.addWidget(self.progress_label)

        row = QHBoxLayout()
        row.setSpacing(8)
        self.kpis = {
            "found": KpiTile("PDFs", ""),
            "added": KpiTile("Added", "success"),
            "duplicates": KpiTile("Duplicates", "skipped"),
            "failed": KpiTile("Failed", "error"),
        }
        for tile in self.kpis.values():
            row.addWidget(tile, 1)
        layout.addLayout(row)
        return card

    def _build_activity_card(self):
        card, layout = make_card()
        top = QHBoxLayout()
        top.addWidget(make_label("Activity", "cardTitle"))
        top.addStretch()
        top.addWidget(make_button("Clear", "link", self.clear_log))
        layout.addLayout(top)

        table = QTableWidget(0, 3)
        table.setHorizontalHeaderLabels(["Time", "Status", "Details"])
        table.verticalHeader().setVisible(False)
        table.verticalHeader().setDefaultSectionSize(30)
        table.setShowGrid(False)
        table.setAlternatingRowColors(True)
        table.setWordWrap(False)
        table.setTextElideMode(Qt.ElideMiddle)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setFocusPolicy(Qt.NoFocus)
        header = table.horizontalHeader()
        header.setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        header.setHighlightSections(False)
        header.setSectionResizeMode(0, QHeaderView.Fixed)
        header.setSectionResizeMode(1, QHeaderView.Fixed)
        header.setSectionResizeMode(2, QHeaderView.Stretch)
        table.setColumnWidth(0, 86)
        table.setColumnWidth(1, 96)
        self.log_table = table
        layout.addWidget(table, 1)
        return card

    def _build_status_bar(self):
        bar = QStatusBar()
        bar.setSizeGripEnabled(False)
        bar.setContentsMargins(12, 2, 12, 2)
        self.setStatusBar(bar)
        self.status_label = make_label("Ready", "muted")
        bar.addWidget(self.status_label, 1)
        self.open_folder_button = make_button("Open Output Folder", slot=self.open_output_folder)
        self.open_workbook_button = make_button("Open Workbook", slot=self.open_output_file)
        for button in (self.open_folder_button, self.open_workbook_button):
            button.setEnabled(False)
            bar.addPermanentWidget(button)

    # ---------------- Activity log ----------------

    def add_log(self, level, message):
        label, color = LEVEL_STYLES.get(level, LEVEL_STYLES["info"])
        row = self.log_table.rowCount()
        self.log_table.insertRow(row)

        time_item = QTableWidgetItem(datetime.now().strftime("%H:%M:%S"))
        time_item.setForeground(QColor(TEXT_SECONDARY))
        status_item = QTableWidgetItem(f"●  {label}")
        status_item.setForeground(QColor(color))
        font = status_item.font()
        font.setBold(True)
        status_item.setFont(font)
        detail_item = QTableWidgetItem(message)
        detail_item.setToolTip(message)

        for column, item in enumerate((time_item, status_item, detail_item)):
            self.log_table.setItem(row, column, item)
        self.log_table.scrollToBottom()

    def clear_log(self):
        self.log_table.setRowCount(0)

    # ---------------- Actions ----------------

    def show_about(self):
        if self._about is None:
            self._about = AboutDialog(self)
        self._about.exec()

    def browse_folder(self):
        current = self.folder_edit.text().strip()
        folder = QFileDialog.getExistingDirectory(
            self, "Select folder with PO PDFs",
            current if os.path.isdir(current) else os.path.expanduser("~"))
        if folder:
            self.folder_edit.setText(os.path.normpath(folder))

    def browse_excel(self):
        current = self.excel_edit.text().strip() or DEFAULT_OUTPUT_FILE
        if not os.path.isdir(os.path.dirname(current)):
            current = os.path.join(os.path.expanduser("~"), os.path.basename(current))
        # Existing workbooks are appended to, never overwritten, so don't
        # ask the user to confirm "replacing" one.
        path, _ = QFileDialog.getSaveFileName(
            self, "Choose output workbook (new or existing)", current,
            "Excel workbook (*.xlsx)", options=QFileDialog.DontConfirmOverwrite)
        if path:
            self.excel_edit.setText(os.path.normpath(path))

    def _open_path(self, path):
        try:
            if sys.platform.startswith("win"):
                os.startfile(path)
            elif sys.platform == "darwin":
                subprocess.run(["open", path])
            else:
                subprocess.run(["xdg-open", path])
        except Exception as e:
            QMessageBox.critical(self, "Couldn't open", str(e))

    def open_output_folder(self):
        if self.last_output_path:
            self._open_path(os.path.dirname(self.last_output_path))

    def open_output_file(self):
        if self.last_output_path and os.path.exists(self.last_output_path):
            self._open_path(self.last_output_path)

    def _set_running(self, running):
        self._running = running
        self.run_button.setEnabled(not running)
        self.run_button.setText("Processing…" if running else "Run Extraction")
        for widget in (self.folder_edit, self.excel_edit):
            widget.setEnabled(not running)
        has_output = bool(self.last_output_path) and not running
        self.open_folder_button.setEnabled(has_output)
        self.open_workbook_button.setEnabled(has_output
                                             and os.path.exists(self.last_output_path))

    def run_extraction(self):
        folder = self.folder_edit.text().strip()
        if not folder or not os.path.isdir(folder):
            QMessageBox.warning(self, "Input folder not found",
                                "Please choose a valid input folder first.")
            return
        output_path = self.excel_edit.text().strip() or DEFAULT_OUTPUT_FILE
        if not output_path.lower().endswith(".xlsx"):
            output_path += ".xlsx"
        output_path = os.path.abspath(output_path)
        self.excel_edit.setText(output_path)
        try:
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
        except OSError as e:
            QMessageBox.warning(self, "Output folder not available", str(e))
            return

        for tile in self.kpis.values():
            tile.set_value(None)
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress_label.setText("Starting…")
        self.status_label.setText("Processing…")
        self._set_running(True)
        threading.Thread(target=self._extract_worker, args=(folder, output_path),
                         daemon=True).start()

    def on_progress(self, done, total):
        self.progress.setRange(0, max(total, 1))
        self.progress.setValue(done)
        self.progress_label.setText(f"Processed {done} of {total} file(s)")

    def on_finished(self, summary):
        if summary.get("output_path"):
            self.last_output_path = summary["output_path"]
        for key in ("found", "added", "duplicates", "failed"):
            if key in summary:
                self.kpis[key].set_value(summary[key])
        self.progress_label.setText(summary["status"])
        self.status_label.setText(summary["status"])
        self._set_running(False)

    # ---------------- Worker (background thread) ----------------

    def _extract_worker(self, folder, output_path):
        log, s = self.signals.log.emit, self.signals

        pdf_files = sorted(
            f for f in os.listdir(folder)
            if f.lower().endswith(".pdf") and os.path.isfile(os.path.join(folder, f))
        )
        if not pdf_files:
            log("warning", f"No PDF files found in {folder}")
            s.finished.emit({"status": "No PDFs found.", "found": 0})
            return
        log("info", f"Found {len(pdf_files)} PDF(s) in {folder}")
        s.progress.emit(0, len(pdf_files))

        # Close the workbook in Excel first, so we can read and save it.
        try:
            message = close_workbook_in_excel(output_path)
            if message:
                log("info", message)
        except Exception as e:
            log("error", str(e))
            s.finished.emit({"status": "Output workbook is locked.",
                             "found": len(pdf_files)})
            return

        pdf_results = []
        failed = 0
        for index, filename in enumerate(pdf_files, start=1):
            pdf_path = os.path.join(folder, filename)
            try:
                rows = extract_fields_from_pdf(pdf_path)
                if rows:
                    pdf_results.append((filename, rows))
                    log("success", f"Read {filename} ({len(rows)} item rows)")
                else:
                    failed += 1
                    log("warning", f"No item lines matched in {filename}")
            except Exception as e:
                failed += 1
                log("error", f"{filename}: {e}")
            s.progress.emit(index, len(pdf_files))

        summary = {"found": len(pdf_files), "added": 0, "duplicates": 0, "failed": failed}
        if not pdf_results:
            log("error", "Nothing extracted. The workbook was not changed.")
            s.finished.emit(dict(summary, status="Nothing extracted."))
            return

        existed_before = os.path.exists(output_path)
        try:
            written, duplicates = build_or_append_workbook(output_path, pdf_results)
        except Exception as e:
            log("error", f"Could not save the workbook: {e}")
            s.finished.emit(dict(summary, status="Save failed."))
            return

        for filename, rows in duplicates:
            log("skipped", f"{filename}: PO No {rows[0]['PO No']} is already in the workbook")

        # Archive every PDF whose data is now in the workbook (including
        # duplicates, since their data was already there).
        moved = 0
        for filename, _rows in written + duplicates:
            try:
                move_to_processed(folder, filename)
                moved += 1
            except Exception as e:
                log("warning", f"Could not move {filename} to "
                               f"'{PROCESSED_FOLDER_NAME}': {e}")

        if written:
            total_rows = sum(len(rows) for _f, rows in written)
            verb = "Appended" if existed_before else "Saved"
            log("success", f"{verb} {total_rows} item rows from {len(written)} PO(s) "
                           f"to {output_path}")
        else:
            log("info", "No new POs to add. The workbook was not changed.")
        if moved:
            log("info", f"{moved} PDF(s) moved to '{PROCESSED_FOLDER_NAME}'.")

        status = (f"Completed with {failed} file(s) needing attention." if failed
                  else "Completed successfully.")
        s.finished.emit(dict(summary, status=status, added=len(written),
                             duplicates=len(duplicates),
                             output_path=output_path if written else None))


def main():
    set_app_user_model_id()
    app = QApplication(sys.argv)
    try:  # keep the light, corporate look even when Windows is in dark mode
        app.styleHints().setColorScheme(Qt.ColorScheme.Light)
    except AttributeError:
        pass
    app.setStyle("Fusion")
    app.setFont(QFont("Segoe UI", 9))
    app.setWindowIcon(QIcon(resource_path("assets", "app.ico")))
    app.setStyleSheet(STYLESHEET)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
